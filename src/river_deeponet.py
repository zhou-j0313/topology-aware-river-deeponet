"""Topology-aware DeepONet baseline for unsteady flow in river networks.

The public release preserves the data pipeline, topology construction,
DeepONet architecture, training loop, validation, and prediction interfaces.
It uses masked mean-squared errors and a simple first-order consistency
regularizer. The paper-specific composite loss and weighting strategy are not
included; the peak and junction soft-loss hooks remain disabled by default.
"""

from __future__ import annotations

import os
import heapq
import time
import warnings
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
import torch.nn.functional as F
from torch.optim import AdamW
from torch.utils.data import DataLoader as TorchDataLoader
from torch.utils.data import Dataset

warnings.filterwarnings("ignore")

torch.manual_seed(42)
np.random.seed(42)
if torch.cuda.is_available():
    torch.cuda.manual_seed_all(42)


# ============================================================
# Standard CSV data loading
# ============================================================

class DataLoader:
    def __init__(self):
        self.cross_sections: Dict[int, Dict[str, np.ndarray]] = {}
        self.river_network: Dict[str, pd.DataFrame] = {}
        self.boundary_conditions: Optional[pd.DataFrame] = None
        self.section_manning_n: Dict[int, float] = {}
        self.upper_boundary_sections: List[int] = []
        self.lower_boundary_sections: List[int] = []

    @staticmethod
    def _check_csv(file_path: str) -> str:
        """Validate a CSV input path."""
        path = str(file_path)
        if not path.lower().endswith(".csv"):
            raise ValueError(f"Only CSV input is supported: {path}")
        if not os.path.isfile(path):
            raise FileNotFoundError(path)
        return path

    def load_cross_sections(self, file_path: str):
        print(f"Loading cross sections: {file_path}", flush=True)
        data_path = self._check_csv(file_path)
        df = pd.read_csv(data_path)
        df.columns = ["section_id", "station_x", "elevation_z"]
        try:
            df["section_id"] = df["section_id"].astype(int)
        except ValueError:
            pass
        self.cross_sections.clear()
        for sec_id, group in df.groupby("section_id"):
            sec_id = int(sec_id)
            self.cross_sections[sec_id] = {
                "X": group["station_x"].values.astype(float),
                "Z": group["elevation_z"].values.astype(float),
                "Z_min": float(group["elevation_z"].min()),
                "Z_max": float(group["elevation_z"].max()),
            }
        print(f"   Loaded {len(self.cross_sections)} cross sections", flush=True)

    def load_river_network(self, file_path: str):
        """Load a headered reach/section network table."""
        print(f"Loading river network: {file_path}", flush=True)
        data_path = self._check_csv(file_path)
        df = pd.read_csv(data_path)
        required = {
            "reach_name", "section_id", "distance_to_next_m", "manning_n"}
        missing = required.difference(df.columns)
        if missing:
            raise ValueError(
                f"River-network CSV is missing columns: {sorted(missing)}")
        self.river_network.clear()
        for reach_name, group in df.groupby("reach_name", sort=False):
            reach_df = group[[
                "section_id", "distance_to_next_m", "manning_n"]].copy()
            reach_df.columns = ["section_id", "distance", "roughness"]
            reach_df["section_id"] = pd.to_numeric(
                reach_df["section_id"], errors="raise").astype(int)
            self.river_network[str(reach_name)] = reach_df.reset_index(drop=True)

        self.section_manning_n.clear()
        for reach_name, reach_df in self.river_network.items():
            if len(reach_df) < 2:
                raise ValueError(
                    f"Reach {reach_name} must contain at least two sections")
            for _, row in reach_df.iterrows():
                self.section_manning_n[int(row["section_id"])] = float(
                    row["roughness"])

        print(f"   Loaded {len(self.river_network)} reaches", flush=True)
        for reach_name, reach_df in self.river_network.items():
            print(f"   - {reach_name}: {len(reach_df)} sections", flush=True)

    def load_boundary_conditions(self, file_path: str):
        print(f"Loading boundary-section configuration: {file_path}", flush=True)
        data_path = self._check_csv(file_path)
        self.boundary_conditions = pd.read_csv(data_path)
        self.boundary_conditions.columns = [
            "upper_boundary_section", "lower_boundary_section"]
        self.boundary_conditions["upper_boundary_section"] = pd.to_numeric(
            self.boundary_conditions["upper_boundary_section"], errors="coerce"
        ).fillna(-1).astype(int)
        self.boundary_conditions["lower_boundary_section"] = pd.to_numeric(
            self.boundary_conditions["lower_boundary_section"], errors="coerce"
        ).fillna(-1).astype(int)
        self.upper_boundary_sections = [
            x for x in self.boundary_conditions[
                "upper_boundary_section"].unique().tolist() if x != -1]
        self.lower_boundary_sections = [
            x for x in self.boundary_conditions[
                "lower_boundary_section"].unique().tolist() if x != -1]
        print(
            f"   Upstream: {self.upper_boundary_sections}; "
            f"downstream: {self.lower_boundary_sections}", flush=True)


# ============================================================
# Hydraulic geometry and lookup tables
# ============================================================

class HydraulicCalculator:
    """
    Compute hydraulic properties for irregular cross sections.

    Water levels above the surveyed banks use either vertical-wall extension,
    sloped overbank extension, or no extension. Surveyed floodplain geometry
    should be preferred whenever it is available.
    """
    def __init__(self, cross_section_data, section_manning_n,
                 overbank_mode: str = "vertical",
                 overbank_side_slope: float = 2.0,
                 overbank_height: float = 20.0):
        self.cross_section_data = cross_section_data
        self.section_manning_n = section_manning_n
        self.overbank_mode = str(overbank_mode).lower()
        if self.overbank_mode not in ["slope", "vertical", "none"]:
            raise ValueError(
                "overbank_mode must be 'slope', 'vertical', or 'none'")
        self.overbank_side_slope = max(float(overbank_side_slope), 0.0)
        self.overbank_height = max(float(overbank_height), 0.0)

    @staticmethod
    def _panel_geometry(x1, z1, x2, z2, wl):
        if z1 >= wl and z2 >= wl:
            return 0.0, 0.0, 0.0
        dx = abs(x2 - x1)
        if dx < 1e-10:
            return 0.0, 0.0, 0.0
        if z1 < wl and z2 < wl:
            h1, h2 = wl - z1, wl - z2
            area = (h1 + h2) / 2.0 * dx
            return area, dx, np.sqrt(dx**2 + (z2 - z1) ** 2)
        if z1 < wl:
            xa, za, xb, zb = x1, z1, x2, z2
        else:
            xa, za, xb, zb = x2, z2, x1, z1
        frac = (wl - za) / (zb - za + 1e-15)
        dx_wet = abs(frac * (xb - xa))
        h_wet = wl - za
        return h_wet * dx_wet / 2.0, dx_wet, np.sqrt(dx_wet**2 + h_wet**2)

    def _extended_geometry_for_wl(self, X, Z, wl):
        """Extend geometry to the requested level and return extra perimeter."""
        X = np.asarray(X, dtype=float)
        Z = np.asarray(Z, dtype=float)
        if X.size < 2:
            return X, Z, 0.0

        # Sort stations to tolerate reversed or unordered input points.
        order = np.argsort(X)
        X = X[order]
        Z = Z[order]

        z_left, z_right = float(Z[0]), float(Z[-1])
        bank_top = max(z_left, z_right, float(np.max(Z)))
        if wl <= bank_top + 1e-10 or self.overbank_mode == "none":
            return X, Z, 0.0

        if self.overbank_mode == "vertical":
            # Vertical walls add perimeter without increasing top width.
            extra_wp = max(0.0, wl - z_left) + max(0.0, wl - z_right)
            return X, Z, extra_wp

        # Slope is horizontal distance per unit vertical rise.
        m = max(self.overbank_side_slope, 1e-6)
        left_dx = max(0.0, wl - z_left) * m
        right_dx = max(0.0, wl - z_right) * m
        X_ext = np.concatenate([[X[0] - left_dx], X, [X[-1] + right_dx]])
        Z_ext = np.concatenate([[wl], Z, [wl]])
        return X_ext, Z_ext, 0.0

    def compute_ABK(self, section_id, wl):
        data = self.cross_section_data[section_id]
        wl = max(float(wl), float(data["Z_min"]) + 1e-6)
        X, Z = data["X"], data["Z"]
        X, Z, extra_wp = self._extended_geometry_for_wl(X, Z, wl)

        area, width, wp = 0.0, 0.0, float(extra_wp)
        for i in range(len(X) - 1):
            a, w, p = self._panel_geometry(X[i], Z[i], X[i + 1], Z[i + 1], wl)
            area += a
            width += w
            wp += p
        area = max(area, 0.001)
        width = max(width, 0.001)
        wp = max(wp, 0.1)
        R = area / wp
        n = self.section_manning_n.get(section_id, 0.03)
        K = (1.0 / n) * area * (R ** (2 / 3))
        return area, width, K

    def geometry_upper_level(self, section_id):
        data = self.cross_section_data[section_id]
        return float(data["Z_max"]) + float(self.overbank_height)

class VectorizedLookupTable:
    def __init__(self, device='cuda'):
        self.device = torch.device(device if torch.cuda.is_available() else 'cpu')
        self.x_locations = []
        self.water_levels_list = []
        self.A_list, self.B_list, self.K_list = [], [], []
        self.n_levels = 41
        self.built = False

    def add_section(self, x_location, layered_ABK):
        self.x_locations.append(float(x_location))
        self.water_levels_list.append(layered_ABK['water_levels'])
        self.A_list.append(layered_ABK['A'])
        self.B_list.append(layered_ABK['B'])
        self.K_list.append(layered_ABK['K'])

    def build(self):
        n_sec = len(self.x_locations)
        self.wl_tensor = torch.zeros(n_sec, self.n_levels, device=self.device)
        self.A_tensor = torch.zeros(n_sec, self.n_levels, device=self.device)
        self.B_tensor = torch.zeros(n_sec, self.n_levels, device=self.device)
        self.K_tensor = torch.zeros(n_sec, self.n_levels, device=self.device)
        for i in range(n_sec):
            n = len(self.water_levels_list[i])
            self.wl_tensor[i, :n] = torch.tensor(self.water_levels_list[i], device=self.device)
            self.A_tensor[i, :n] = torch.tensor(self.A_list[i], device=self.device)
            self.B_tensor[i, :n] = torch.tensor(self.B_list[i], device=self.device)
            self.K_tensor[i, :n] = torch.tensor(self.K_list[i], device=self.device)
            if n < self.n_levels:
                for t in [self.wl_tensor, self.A_tensor, self.B_tensor, self.K_tensor]:
                    t[i, n:] = t[i, n - 1]
        self.z_min_tensor = self.wl_tensor[:, 0]
        self.z_max_tensor = self.wl_tensor[:, -1]
        self.built = True
        print(
            f"Hydraulic lookup table: {n_sec} sections, "
            f"{self.n_levels} levels per section", flush=True)

    def query_batch(self, x_indices, water_levels):
        """Differentiably interpolate area, top width, and conveyance."""
        if not self.built:
            self.build()
        wl = water_levels.squeeze(-1) if water_levels.dim() > 1 else water_levels
        wl_sections = self.wl_tensor[x_indices]
        A_sections = self.A_tensor[x_indices]
        B_sections = self.B_tensor[x_indices]
        K_sections = self.K_tensor[x_indices]
        z_min = self.z_min_tensor[x_indices]
        z_max = self.z_max_tensor[x_indices]
        wl_clamped = torch.clamp(wl, min=z_min, max=z_max)
        wl_normalized = (wl_clamped - z_min) / (z_max - z_min + 1e-8) * (self.n_levels - 1)
        idx_lower = torch.floor(wl_normalized).long().clamp(0, self.n_levels - 2)
        idx_upper = idx_lower + 1
        w = wl_normalized - idx_lower.float()
        batch_idx = torch.arange(len(x_indices), device=self.device)
        A = A_sections[batch_idx, idx_lower] * (1 - w) + A_sections[batch_idx, idx_upper] * w
        B = B_sections[batch_idx, idx_lower] * (1 - w) + B_sections[batch_idx, idx_upper] * w
        K = K_sections[batch_idx, idx_lower] * (1 - w) + K_sections[batch_idx, idx_upper] * w
        return A, B, K


# ============================================================
# River-network event data structures
# ============================================================

@dataclass
class EventSpec:
    name: str
    boundary_file: str
    result_file: Optional[str] = None
    lateral_inflow_file: Optional[str] = None
    initial_condition_file: Optional[str] = None
    observed_sections: Optional[List[object]] = None


@dataclass
class RawEvent:
    """Raw time series for one event."""
    name: str
    t_hours: np.ndarray

    # Multiple boundary series used as model inputs.
    upper_q: Dict[int, np.ndarray] = field(default_factory=dict)
    lower_z: Dict[int, np.ndarray] = field(default_factory=dict)

    # Repeated junction endpoints occupy separate reach-section rows.
    lateral: np.ndarray = field(default_factory=lambda: np.zeros((0, 0)))
    node_source: np.ndarray = field(default_factory=lambda: np.zeros((0, 0)))
    z0: np.ndarray = field(default_factory=lambda: np.zeros(0))
    q0: np.ndarray = field(default_factory=lambda: np.zeros(0))

    # Supervision targets used only by the training data loss.
    target_z: Optional[np.ndarray] = None
    target_q: Optional[np.ndarray] = None

    # Reference targets are used for evaluation, never as model inputs.
    reference_z: Optional[np.ndarray] = None
    reference_q: Optional[np.ndarray] = None

    observed_section_indices: Optional[List[int]] = None
    reference_section_indices: Optional[List[int]] = None


@dataclass
class PreparedEvent:
    name: str
    t_hours_raw: np.ndarray
    duration_hours: float
    input_grid: torch.Tensor

    # Training targets and observation mask.
    target_grid: Optional[torch.Tensor]
    obs_mask: torch.Tensor

    # Evaluation targets and reference mask.
    reference_grid: Optional[torch.Tensor]
    reference_mask: torch.Tensor

    time_mask: torch.Tensor
    valid_T: int
    dt_hours: float

class EventTensorDataset(Dataset):
    def __init__(self, events: List[PreparedEvent], indices: List[int]):
        self.events = events
        self.indices = indices

    def __len__(self):
        return len(self.indices)

    def __getitem__(self, idx):
        e = self.events[self.indices[idx]]
        N, T = e.input_grid.shape[1], e.input_grid.shape[2]
        if e.target_grid is not None:
            y = e.target_grid
        else:
            y = torch.zeros(2, N, T, device=e.input_grid.device)

        if e.reference_grid is not None:
            ref_y = e.reference_grid
        else:
            ref_y = torch.zeros(2, N, T, device=e.input_grid.device)

        return {
            'x': e.input_grid,
            'y': y,
            'mask': e.obs_mask,
            'ref_y': ref_y,
            'ref_mask': e.reference_mask,
            'time_mask': e.time_mask,
            'idx': self.indices[idx],
        }


# ============================================================
# Neural-network modules
# ============================================================

class TemporalResidualBlock(nn.Module):
    """Dilated temporal residual block preserving time resolution."""

    def __init__(self, channels, dilation):
        super().__init__()
        padding = 2 * int(dilation)
        self.depthwise = nn.Conv2d(
            channels, channels,
            kernel_size=(1, 5),
            dilation=(1, int(dilation)),
            padding=(0, padding),
            groups=channels,
        )
        self.mix = nn.Sequential(
            nn.Conv2d(channels, channels * 2, kernel_size=1),
            nn.GELU(),
            nn.Conv2d(channels * 2, channels, kernel_size=1),
        )
        self.norm = nn.GroupNorm(1, channels)

    def forward(self, x):
        return self.norm(x + self.mix(self.depthwise(x)))


class HydraulicGraphMessageBlock(nn.Module):
    """
    Directed hydraulic graph message-passing layer.

    Reach edges use physical distances and learnable travel speeds. Forward
    edges represent downstream propagation, reverse edges represent backwater
    influence, and junction edges have zero distance.
    """

    def __init__(
            self, hidden, edge_dim, edge_index, edge_attr, edge_dx_m,
            dropout=0.05):
        super().__init__()
        self.hidden = int(hidden)
        self.edge_dim = int(edge_dim)
        self.register_buffer(
            'edge_index', edge_index.long(), persistent=False)
        self.register_buffer(
            'edge_attr', edge_attr.float(), persistent=False)
        self.register_buffer(
            'edge_dx_m', edge_dx_m.float(), persistent=False)

        self.message = nn.Sequential(
            nn.Linear(2 * hidden + edge_dim, hidden * 2),
            nn.GELU(),
            nn.Dropout(dropout),
            nn.Linear(hidden * 2, hidden),
        )
        self.gate = nn.Sequential(
            nn.Linear(2 * hidden + edge_dim, hidden),
            nn.Sigmoid(),
        )
        self.update = nn.Sequential(
            nn.Linear(2 * hidden, hidden * 2),
            nn.GELU(),
            nn.Dropout(dropout),
            nn.Linear(hidden * 2, hidden),
        )
        self.norm = nn.GroupNorm(1, hidden)

        # Initial downstream and backwater propagation-speed parameters.
        self.log_downstream_speed = nn.Parameter(torch.tensor(1.12))
        self.log_backwater_speed = nn.Parameter(torch.tensor(3.88))

    @staticmethod
    def _positive_speed(parameter):
        return F.softplus(parameter) + 0.1

    def _sample_delayed_source(self, node_features, dt_hours):
        # node_features: [B,H,N,T]
        batch, hidden, _, n_times = node_features.shape
        src = self.edge_index[0]
        source_history = node_features.index_select(2, src)
        # [B,H,E,T] -> [B,E,T,H]
        source_history = source_history.permute(0, 2, 3, 1).contiguous()

        direction = self.edge_attr[:, 3]
        downstream_speed = self._positive_speed(
            self.log_downstream_speed)
        backwater_speed = self._positive_speed(
            self.log_backwater_speed)
        speed = torch.where(
            direction >= 0.0, downstream_speed, backwater_speed)

        dt_seconds = dt_hours.to(
            node_features.dtype).clamp(min=1e-6) * 3600.0
        lag = self.edge_dx_m.view(1, -1) / (
            speed.view(1, -1) * dt_seconds.view(-1, 1))
        lag = lag.clamp(min=0.0, max=float(max(n_times - 1, 0)))

        time = torch.arange(
            n_times, dtype=node_features.dtype,
            device=node_features.device).view(1, 1, -1)
        sample_time = time - lag.unsqueeze(-1)
        valid = sample_time >= 0.0
        sample_time = sample_time.clamp(
            min=0.0, max=float(max(n_times - 1, 0)))
        lower = torch.floor(sample_time).long()
        upper = torch.clamp(lower + 1, max=max(n_times - 1, 0))
        alpha = (sample_time - lower.to(node_features.dtype)).unsqueeze(-1)

        index_lower = lower.unsqueeze(-1).expand(
            batch, -1, -1, hidden)
        index_upper = upper.unsqueeze(-1).expand(
            batch, -1, -1, hidden)
        value_lower = torch.gather(
            source_history, 2, index_lower)
        value_upper = torch.gather(
            source_history, 2, index_upper)
        delayed = value_lower * (1.0 - alpha) + value_upper * alpha
        initial = source_history[:, :, :1, :].expand_as(delayed)
        return torch.where(valid.unsqueeze(-1), delayed, initial)

    def forward(self, node_features, dt_hours):
        if self.edge_index.numel() == 0:
            return node_features

        batch, hidden, n_nodes, n_times = node_features.shape
        src = self.edge_index[0]
        dst = self.edge_index[1]
        delayed_source = self._sample_delayed_source(
            node_features, dt_hours)
        destination = node_features.index_select(
            2, dst).permute(0, 2, 3, 1).contiguous()
        edge_features = self.edge_attr.view(
            1, -1, 1, self.edge_dim).expand(
            batch, -1, n_times, -1)

        message_input = torch.cat(
            [delayed_source, destination, edge_features], dim=-1)
        message = self.message(message_input) * self.gate(message_input)

        aggregate = torch.zeros(
            batch, n_nodes, n_times, hidden,
            dtype=message.dtype, device=message.device)
        aggregate.index_add_(1, dst, message)
        degree = torch.zeros(
            n_nodes, dtype=message.dtype, device=message.device)
        degree.index_add_(
            0, dst, torch.ones_like(dst, dtype=message.dtype))
        aggregate = aggregate / degree.clamp(
            min=1.0).view(1, n_nodes, 1, 1)

        current = node_features.permute(0, 2, 3, 1)
        delta = self.update(torch.cat([current, aggregate], dim=-1))
        updated = (current + delta).permute(0, 3, 1, 2).contiguous()
        return self.norm(updated)


class HydraulicGraphStack(nn.Module):
    """Stack of directed hydraulic graph message-passing blocks."""

    def __init__(
            self, hidden, edge_dim, edge_index, edge_attr, edge_dx_m,
            layers=2):
        super().__init__()
        self.blocks = nn.ModuleList([
            HydraulicGraphMessageBlock(
                hidden=hidden,
                edge_dim=edge_dim,
                edge_index=edge_index,
                edge_attr=edge_attr,
                edge_dx_m=edge_dx_m,
            )
            for _ in range(max(int(layers), 1))
        ])

    def forward(self, x, dt_hours):
        for block in self.blocks:
            x = block(x, dt_hours)
        return x


class TopologyAwareBranch(nn.Module):
    """
    Branch encoder preserving temporal phase, reach locality, and boundary paths.

    It returns node- and time-specific DeepONet coefficients instead of one
    static latent vector per event.
    """

    def __init__(
            self, in_channels, hidden, latent,
            reach_membership, junction_mask,
            boundary_channel_positions, boundary_route_weights,
            boundary_lag_fraction, dt_channel_position,
            edge_index, edge_attr, edge_dx_m):
        super().__init__()
        self.hidden = int(hidden)
        self.latent = int(latent)
        self.register_buffer(
            'reach_membership', reach_membership.float(), persistent=False)
        self.register_buffer(
            'junction_mask', junction_mask.float(), persistent=False)
        self.register_buffer(
            'boundary_channel_positions',
            boundary_channel_positions.long(), persistent=False)
        self.register_buffer(
            'boundary_route_weights',
            boundary_route_weights.float(), persistent=False)
        self.register_buffer(
            'boundary_lag_fraction',
            boundary_lag_fraction.float(), persistent=False)
        self.dt_channel_position = int(dt_channel_position)
        # Learn one maximum propagation lag per boundary series.
        self.boundary_lag_logits = nn.Parameter(torch.full(
            (int(boundary_route_weights.shape[1]),), -1.1))

        self.point_encoder = nn.Sequential(
            nn.Conv2d(in_channels, hidden, kernel_size=1),
            nn.GELU(),
            nn.Conv2d(hidden, hidden, kernel_size=1),
            nn.GELU(),
        )
        self.temporal_encoder = nn.Sequential(
            TemporalResidualBlock(hidden, dilation=1),
            TemporalResidualBlock(hidden, dilation=2),
            TemporalResidualBlock(hidden, dilation=4),
            TemporalResidualBlock(hidden, dilation=8),
        )
        self.graph_encoder = HydraulicGraphStack(
            hidden=hidden,
            edge_dim=int(edge_attr.shape[1]),
            edge_index=edge_index,
            edge_attr=edge_attr,
            edge_dx_m=edge_dx_m,
            layers=2,
        )

        self.boundary_encoder = nn.Sequential(
            nn.Conv1d(1, hidden, kernel_size=5, padding=2),
            nn.GELU(),
            nn.Conv1d(hidden, hidden, kernel_size=5, padding=4, dilation=2),
            nn.GELU(),
            nn.Conv1d(hidden, hidden, kernel_size=5, padding=8, dilation=4),
            nn.GELU(),
        )

        # Fuse local, reach, global, junction, and boundary-path context.
        self.fusion = nn.Sequential(
            nn.Conv2d(hidden * 5, hidden * 2, kernel_size=1),
            nn.GELU(),
            TemporalResidualBlock(hidden * 2, dilation=1),
            nn.Conv2d(hidden * 2, 2 * latent + 2, kernel_size=1),
        )

    def _delayed_boundary_context(self, encoded):
        """Differentiably delay boundary features using path distance."""
        batch, n_boundaries, hidden, n_times = encoded.shape
        context = torch.zeros(
            batch, self.boundary_route_weights.shape[0], hidden, n_times,
            dtype=encoded.dtype, device=encoded.device)
        time = torch.arange(
            n_times, dtype=encoded.dtype, device=encoded.device).view(1, -1)
        max_lag = torch.sigmoid(
            self.boundary_lag_logits) * max(n_times - 1, 1)

        for boundary_idx in range(n_boundaries):
            lag = (
                self.boundary_lag_fraction[:, boundary_idx:boundary_idx + 1]
                * max_lag[boundary_idx]
            )
            sample_time = time - lag
            valid = sample_time >= 0.0
            sample_time = sample_time.clamp(0.0, float(max(n_times - 1, 0)))
            lower = torch.floor(sample_time).long()
            upper = torch.clamp(lower + 1, max=max(n_times - 1, 0))
            alpha = (sample_time - lower.to(encoded.dtype)).view(
                1, -1, 1, n_times)

            series = encoded[:, boundary_idx]
            expanded = series[:, None, :, :].expand(
                batch, lag.shape[0], hidden, n_times)
            lower_index = lower.view(
                1, -1, 1, n_times).expand(batch, -1, hidden, -1)
            upper_index = upper.view(
                1, -1, 1, n_times).expand(batch, -1, hidden, -1)
            sampled_lower = torch.gather(expanded, 3, lower_index)
            sampled_upper = torch.gather(expanded, 3, upper_index)
            sampled = (
                sampled_lower * (1.0 - alpha)
                + sampled_upper * alpha
            )
            initial = expanded[..., :1].expand_as(sampled)
            sampled = torch.where(
                valid.view(1, -1, 1, n_times), sampled, initial)
            weight = self.boundary_route_weights[
                :, boundary_idx].view(1, -1, 1, 1)
            context = context + sampled * weight
        return context.permute(0, 2, 1, 3).contiguous()

    def forward(self, x):
        # x: [B,C,N,T]
        batch, _, n_nodes, n_times = x.shape
        local = self.temporal_encoder(self.point_encoder(x))
        dt_hours = x[:, self.dt_channel_position, 0, 0]
        local = self.graph_encoder(local, dt_hours)

        # Pool each reach independently while preserving every time step.
        reach_weights = self.reach_membership
        reach_sum = torch.einsum(
            'bcnt,rn->bcrt', local, reach_weights)
        reach_den = reach_weights.sum(dim=1).clamp(min=1.0).view(
            1, 1, -1, 1)
        reach_features = reach_sum / reach_den
        reach_context = torch.einsum(
            'bcrt,rn->bcnt', reach_features, reach_weights)

        global_context = local.mean(
            dim=2, keepdim=True).expand(-1, -1, n_nodes, -1)

        junction_den = self.junction_mask.sum().clamp(min=1.0)
        junction_feature = (
            local * self.junction_mask.view(1, 1, -1, 1)
        ).sum(dim=2, keepdim=True) / junction_den
        junction_context = junction_feature.expand(
            -1, -1, n_nodes, -1)

        # Encode each boundary series and route it to reachable nodes.
        boundary_series = x.index_select(
            1, self.boundary_channel_positions).mean(dim=2)
        n_boundaries = int(boundary_series.shape[1])
        boundary_encoded = self.boundary_encoder(
            boundary_series.reshape(batch * n_boundaries, 1, n_times)
        ).reshape(batch, n_boundaries, self.hidden, n_times)
        boundary_context = self._delayed_boundary_context(boundary_encoded)

        fused = self.fusion(torch.cat([
            local,
            reach_context,
            global_context,
            junction_context,
            boundary_context,
        ], dim=1))
        fused = fused.permute(0, 2, 3, 1)
        branch_h = fused[..., :self.latent]
        branch_q = fused[..., self.latent:2 * self.latent]
        bias_h = fused[..., 2 * self.latent]
        bias_q = fused[..., 2 * self.latent + 1]
        return branch_h, branch_q, bias_h, bias_q


class TopologyAwareTrunk(nn.Module):
    """Trunk network for continuous coordinates and topology features."""

    def __init__(self, input_dim, hidden, latent):
        super().__init__()
        self.net = nn.Sequential(
            nn.Linear(input_dim, hidden),
            nn.GELU(),
            nn.Linear(hidden, hidden),
            nn.GELU(),
            nn.Linear(hidden, hidden),
            nn.GELU(),
            nn.Linear(hidden, latent),
        )

    def forward(self, coords):
        return self.net(coords)


class TopologyDeepONet(nn.Module):
    """Path-conditioned river-network DeepONet with temporal phase."""

    def __init__(
            self, in_channels, query_dim, hidden, latent,
            reach_membership, junction_mask,
            boundary_channel_positions, boundary_route_weights,
            boundary_lag_fraction, dt_channel_position,
            edge_index, edge_attr, edge_dx_m):
        super().__init__()
        self.branch = TopologyAwareBranch(
            in_channels=in_channels,
            hidden=hidden,
            latent=latent,
            reach_membership=reach_membership,
            junction_mask=junction_mask,
            boundary_channel_positions=boundary_channel_positions,
            boundary_route_weights=boundary_route_weights,
            boundary_lag_fraction=boundary_lag_fraction,
            dt_channel_position=dt_channel_position,
            edge_index=edge_index,
            edge_attr=edge_attr,
            edge_dx_m=edge_dx_m,
        )
        self.trunk_h = TopologyAwareTrunk(query_dim, hidden, latent)
        self.trunk_q = TopologyAwareTrunk(query_dim, hidden, latent)

    def forward(self, x, query_coords, n_nodes, n_times):
        branch_h, branch_q, bias_h, bias_q = self.branch(x)
        trunk_h = self.trunk_h(query_coords).reshape(
            n_nodes, n_times, self.branch.latent)
        trunk_q = self.trunk_q(query_coords).reshape(
            n_nodes, n_times, self.branch.latent)
        h = (branch_h * trunk_h.unsqueeze(0)).sum(dim=-1) + bias_h
        q = (branch_q * trunk_q.unsqueeze(0)).sum(dim=-1) + bias_q
        return torch.stack([h, q], dim=1)


# ============================================================
# Training history
# ============================================================

@dataclass
class TrainingHistory:
    epoch: List[int] = field(default_factory=list)
    lr: List[float] = field(default_factory=list)
    train_loss: List[float] = field(default_factory=list)
    train_h: List[float] = field(default_factory=list)
    train_q: List[float] = field(default_factory=list)
    train_bc: List[float] = field(default_factory=list)
    train_peak: List[float] = field(default_factory=list)
    train_pde_cont: List[float] = field(default_factory=list)
    train_pde_mom: List[float] = field(default_factory=list)
    train_rel: List[float] = field(default_factory=list)
    val_loss: List[float] = field(default_factory=list)
    elapsed_sec: List[float] = field(default_factory=list)

    def to_dataframe(self) -> pd.DataFrame:
        n = len(self.epoch)
        return pd.DataFrame({
            'epoch': self.epoch[:n], 'lr': self.lr[:n],
            'train_loss': self.train_loss[:n],
            'loss_h': self.train_h[:n], 'loss_q': self.train_q[:n],
            'loss_bc': self.train_bc[:n], 'loss_peak': self.train_peak[:n],
            'loss_pde_cont': self.train_pde_cont[:n],
            'loss_pde_mom': self.train_pde_mom[:n],
            'loss_rel': self.train_rel[:n],
            'val_loss': self.val_loss[:n], 'elapsed_s': self.elapsed_sec[:n],
        })


# ============================================================
# River-network neural operator
# ============================================================

class RiverOperatorSurrogate(nn.Module):
    """
    Topology-aware DeepONet surrogate for a branched river network.

    The system expands the network into reach-section rows, identifies repeated
    junction endpoints, supports multiple external boundaries, projects
    predictions onto hard constraints, and evaluates consistency reach by reach.
    """

    def __init__(self, data_loader, n_time_model=None,
                 device='cuda', deeponet_hidden=128, deeponet_latent=128,
                 boundary_hard_injection=True,
                 overbank_mode: str = 'slope',
                 overbank_side_slope: float = 2.0,
                 overbank_height: float = 20.0,
                 **unused_loss_options):
        super().__init__()
        self.dl = data_loader
        self.n_time_model = None if n_time_model is None else int(n_time_model)
        self.device = torch.device(device if torch.cuda.is_available() else 'cpu')
        self.calc = HydraulicCalculator(
            self.dl.cross_sections, self.dl.section_manning_n,
            overbank_mode=overbank_mode,
            overbank_side_slope=overbank_side_slope,
            overbank_height=overbank_height)
        self.lut = VectorizedLookupTable(device=str(self.device))
        self.overbank_mode = str(overbank_mode).lower()
        self.overbank_side_slope = float(overbank_side_slope)
        self.overbank_height = float(overbank_height)
        # Accept legacy loss options for configuration compatibility.
        self.unused_loss_options = dict(unused_loss_options)
        self.boundary_hard_injection = boundary_hard_injection
        self.g = 9.81
        self.deeponet_hidden = int(deeponet_hidden)
        self.deeponet_latent = int(deeponet_latent)

        self.event_specs: List[EventSpec] = []
        self.raw_events: List[RawEvent] = []
        self.events: List[PreparedEvent] = []
        self.train_event_names: List[str] = []
        self.val_event_names: List[str] = []
        self.test_event_names: List[str] = []

        # Junction sections may occur in multiple reach-section rows.
        self.reach_names: List[str] = []
        self.reach_section_ids: Dict[str, List[int]] = {}
        self.reach_row_indices: Dict[str, List[int]] = {}
        self.reach_local_x: Dict[str, np.ndarray] = {}
        self.reach_global_x: Dict[str, np.ndarray] = {}
        self.reach_lengths: Dict[str, float] = {}
        self.row_to_meta: List[Dict] = []
        self.row_sec_ids: List[int] = []
        self.row_z_bed: List[float] = []
        self.row_local_x: List[float] = []
        self.row_global_x: List[float] = []
        self.row_reach_ids: List[int] = []
        self.row_manning_n: List[float] = []
        self.row_lut_indices: List[int] = []
        self.row_indices_by_section: Dict[int, List[int]] = {}
        self.section_representative_row: Dict[int, int] = {}
        self.section_occurrences: Dict[int, List[Dict]] = {}
        self.duplicate_section_ids = set()
        self.junctions: Dict[int, Dict] = {}

        self.section_to_idx: Dict[int, int] = {}
        self.section_z_bed: Dict[int, float] = {}
        self.section_map: Dict[int, float] = {}

        self.x_coords = None
        self.x_local_array = None
        self.x_global_array = None
        self.reach_id_array = None
        self.z_bed_array = None
        self.S0_array = None
        self.manning_n_array = None
        self.L = None
        self.prepared = False

        self.q_center = 0.0
        self.q_scale = 1.0
        self.q_min = 0.0
        self.q_max = 1.0
        self.h_center = 0.0
        self.h_scale = 1.0
        self.h_min = 0.0
        self.h_max = 1.0
        self.source_center = 0.0
        self.source_scale = 1.0
        self.source_min = 0.0
        self.source_max = 1.0
        self.duration_scale = 1.0
        self.log_qmax_center = 0.0
        self.log_qmax_scale = 1.0
        self.log_hrange_center = 0.0
        self.log_hrange_scale = 1.0

        self.z_bed_tensor = None
        self.row_lut_indices_tensor = None
        self.reach_dx_tensor: Dict[str, torch.Tensor] = {}
        self.sec_idx_tensor = None
        self.query_coords = None
        self.branch_channel_indices_tensor = None
        self.model = None

        self.upper_boundary_sections_sorted = sorted([int(x) for x in self.dl.upper_boundary_sections])
        self.lower_boundary_sections_sorted = sorted([int(x) for x in self.dl.lower_boundary_sections])
        self.CHANNELS = self._build_channel_layout()
        self.in_channels = len(self.CHANNELS)

        self.history = TrainingHistory()

        print(f"Device: {self.device}", flush=True)
        print("Framework: topology-aware DeepONet-PINO v8.1", flush=True)
        print(
            f"Boundaries: upstream={self.upper_boundary_sections_sorted}, "
            f"downstream={self.lower_boundary_sections_sorted}", flush=True)
        print(f"Input channels: {self.in_channels}", flush=True)
        print(
            "Boundary inputs: upstream Q and downstream Z/h only", flush=True)

    # ----------------------------------------------------------
    # Input-channel layout
    # ----------------------------------------------------------

    def _build_channel_layout(self) -> Dict[str, int]:
        """
        Keep training and prediction inputs consistent. Only upstream discharge
        and downstream water level/depth are supplied as boundary conditions.

        """
        names = [
            'x_norm', 'x_local_norm', 'x_global_norm', 'reach_norm',
            'z_bed_norm', 'S0_norm', 'n_norm', 'tau', 'time_mask', 'dt_hours',

            # Available boundary channels.
            'q_up', 'h_low',

            # Event-relative boundary channels.
            'q_up_evt', 'h_low_rel',

            # Sources, initial conditions, and event conditioning.
            'lateral', 'lateral_cum', 'node_source',
            'lateral_evt', 'lateral_cum_evt', 'node_source_evt',
            'h0', 'q0', 'q0_evt', 'duration',
            'log_qmax_evt', 'log_hrange_evt',
        ]
        for sec in self.upper_boundary_sections_sorted:
            names.extend([f'up_{sec}_q', f'up_{sec}_q_evt'])
        for sec in self.lower_boundary_sections_sorted:
            names.extend([f'low_{sec}_h', f'low_{sec}_h_rel'])
        for sec in self.upper_boundary_sections_sorted:
            names.append(f'mask_up_{sec}')
        for sec in self.lower_boundary_sections_sorted:
            names.append(f'mask_low_{sec}')
        return {name: i for i, name in enumerate(names)}

    # ----------------------------------------------------------
    # Public interface
    # ----------------------------------------------------------

    def add_event(self, name, boundary_file, result_file=None,
                  lateral_inflow_file=None, initial_condition_file=None):
        self.event_specs.append(EventSpec(
            name=name, boundary_file=boundary_file, result_file=result_file,
            lateral_inflow_file=lateral_inflow_file,
            initial_condition_file=initial_condition_file,
            observed_sections=None))
        print(f"Registered event: {name} (fully supervised)", flush=True)

    def add_event_sparse(self, name, boundary_file, result_file=None,
                         lateral_inflow_file=None, initial_condition_file=None,
                         observed_sections=None):
        self.event_specs.append(EventSpec(
            name=name, boundary_file=boundary_file, result_file=result_file,
            lateral_inflow_file=lateral_inflow_file,
            initial_condition_file=initial_condition_file,
            observed_sections=observed_sections))
        # None uses all targets, [] uses no observations, and a list selects
        # sparse observation sections.
        if observed_sections is None:
            info = "fully supervised" if result_file else "boundary driven"
        elif len(observed_sections) == 0:
            info = "no observation supervision"
            if result_file:
                info += ", reference targets retained"
        else:
            info = f"sparse observations={observed_sections}"
        print(f"Registered event: {name} ({info})", flush=True)

    def set_split(self, train_event_names, val_event_names=None, test_event_names=None):
        self.train_event_names = list(train_event_names)
        self.val_event_names = list(val_event_names or [])
        self.test_event_names = list(test_event_names or [])

    def set_events_observed_sections(self, event_names, observed_sections):
        """
        Update observation sections for registered events.

        An empty list makes validation boundary-driven while retaining target
        files strictly for evaluation.
        """
        name_set = set(event_names or [])
        changed = 0
        for spec in self.event_specs:
            if spec.name in name_set:
                spec.observed_sections = observed_sections
                changed += 1
        if changed > 0:
            print(
                f"Updated observation sections for {changed} events: "
                f"{observed_sections}", flush=True)

    def prepare_data(self):
        print("\n" + "=" * 72, flush=True)
        print("Preparing DeepONet-PINO data and river topology", flush=True)
        print("=" * 72, flush=True)
        self._build_network_topology()
        self._build_static_arrays()
        self._build_physical_topology_features()
        self._build_lookup_table_if_needed()
        self._load_all_events()
        self._set_time_grid_from_events()
        self._compute_normalization()
        self._build_prepared_events()
        self._build_query_coords()
        self._build_deeponet_model()
        self.prepared = True
        if not self.train_event_names:
            self.train_event_names = [e.name for e in self.events]
        n_params = sum(p.numel() for p in self.model.parameters())
        obs_ratios = [e.obs_mask.mean().item() for e in self.events]
        avg_obs = np.mean(obs_ratios) if obs_ratios else 0
        print(
            f"Data prepared: events={len(self.events)}, parameters={n_params:,}",
            flush=True)
        print(
            f"   Network rows={len(self.row_sec_ids)}, physical sections="
            f"{len(self.section_representative_row)}", flush=True)
        print(f"   Mean observation coverage={avg_obs*100:.1f}%", flush=True)

    # ----------------------------------------------------------
    # River topology and automatic splitting at internal connections
    # ----------------------------------------------------------

    def _build_network_topology(self):
        raw_reach_names = list(self.dl.river_network.keys())
        if len(raw_reach_names) == 0:
            raise ValueError(
                "No reaches loaded; call dl.load_river_network() first")

        raw_occurrences = {}
        for reach_name in raw_reach_names:
            raw_df = self.dl.river_network[reach_name].reset_index(drop=True)
            sec_ids = [int(x) for x in raw_df['section_id'].tolist()]
            n_sec = len(sec_ids)
            for j, sec_id in enumerate(sec_ids):
                raw_occurrences.setdefault(sec_id, []).append({
                    'reach_name': reach_name, 'local_idx': j,
                    'is_head': j == 0, 'is_tail': j == n_sec - 1,
                })

        expanded_reaches = []
        split_count = 0
        for reach_name in raw_reach_names:
            reach_df = self.dl.river_network[reach_name].reset_index(drop=True).copy()
            reach_df['section_id'] = reach_df['section_id'].astype(int)
            sec_ids = reach_df['section_id'].tolist()
            n_sec = len(sec_ids)
            split_positions = sorted({
                j for j, sec_id in enumerate(sec_ids)
                if 0 < j < n_sec - 1 and len(raw_occurrences.get(sec_id, [])) > 1
            })
            if len(split_positions) == 0:
                expanded_reaches.append((reach_name, reach_df))
            else:
                split_count += len(split_positions)
                cut_points = [0] + split_positions + [n_sec - 1]
                for k in range(len(cut_points) - 1):
                    s, e = cut_points[k], cut_points[k + 1]
                    seg_df = reach_df.iloc[s:e + 1].reset_index(drop=True).copy()
                    expanded_reaches.append((f"{reach_name}_seg{k + 1}", seg_df))

        self.reach_names = [name for name, _ in expanded_reaches]
        self.reach_section_ids.clear()
        self.reach_row_indices.clear()
        self.reach_local_x.clear()
        self.reach_global_x.clear()
        self.reach_lengths.clear()
        self.row_to_meta.clear()
        self.row_sec_ids.clear()
        self.row_z_bed.clear()
        self.row_local_x.clear()
        self.row_global_x.clear()
        self.row_reach_ids.clear()
        self.row_manning_n.clear()
        self.row_indices_by_section.clear()
        self.section_representative_row.clear()
        self.section_occurrences.clear()
        self.duplicate_section_ids.clear()

        global_cursor = 0.0
        row_id = 0
        for r_idx, (reach_name, reach_df) in enumerate(expanded_reaches):
            sec_ids = [int(x) for x in reach_df['section_id'].tolist()]
            local_x = []
            cur = 0.0
            for i in range(len(sec_ids)):
                local_x.append(cur)
                if i < len(sec_ids) - 1:
                    step = float(reach_df.loc[i, 'distance']) if pd.notna(
                        reach_df.loc[i, 'distance']) else 0.0
                    cur += step
            reach_len = float(max(cur, 1.0))
            self.reach_section_ids[reach_name] = sec_ids
            self.reach_local_x[reach_name] = np.array(local_x, dtype=float)
            self.reach_global_x[reach_name] = global_cursor + np.array(local_x, dtype=float)
            self.reach_lengths[reach_name] = reach_len

            rows = []
            for j, sec_id in enumerate(sec_ids):
                if sec_id not in self.dl.cross_sections:
                    raise KeyError(
                        f"Section {sec_id} is in the network but missing "
                        "from the cross-section file")
                roughness = (
                    float(reach_df.loc[j, 'roughness'])
                    if pd.notna(reach_df.loc[j, 'roughness'])
                    else self.dl.section_manning_n.get(sec_id, 0.03))
                meta = {
                    'row': row_id, 'reach_name': reach_name, 'reach_idx': r_idx,
                    'local_idx': j, 'sec_id': sec_id,
                    'is_head': j == 0, 'is_tail': j == len(sec_ids) - 1,
                    'x_local': float(local_x[j]), 'x_global': float(global_cursor + local_x[j]),
                }
                rows.append(row_id)
                self.row_to_meta.append(meta)
                self.row_sec_ids.append(sec_id)
                self.row_z_bed.append(float(self.dl.cross_sections[sec_id]['Z_min']))
                self.row_local_x.append(meta['x_local'])
                self.row_global_x.append(meta['x_global'])
                self.row_reach_ids.append(r_idx)
                self.row_manning_n.append(roughness)
                self.row_indices_by_section.setdefault(sec_id, []).append(row_id)
                self.section_occurrences.setdefault(sec_id, []).append(meta)
                if sec_id not in self.section_representative_row:
                    self.section_representative_row[sec_id] = row_id
                row_id += 1

            self.reach_row_indices[reach_name] = rows
            # This display coordinate is not used by DeepONet.
            global_cursor += reach_len + 1.0e3

        for sec_id, occs in self.section_occurrences.items():
            if len(occs) > 1:
                self.duplicate_section_ids.add(sec_id)

        self._identify_junctions()

        print(
            f"River topology: {len(raw_reach_names)} input reaches, "
            f"{len(self.reach_names)} reaches after splitting", flush=True)
        if split_count > 0:
            print(f"   Internal connection splits: {split_count}", flush=True)
        print(
            f"   Reach-section rows={len(self.row_sec_ids)}; repeated "
            f"sections={len(self.duplicate_section_ids)}", flush=True)
        print(f"   Junctions identified={len(self.junctions)}", flush=True)

    def _identify_junctions(self):
        self.junctions.clear()
        for sec_id, occs in self.section_occurrences.items():
            endpoint_occs = [m for m in occs if m['is_head'] or m['is_tail']]
            if len(endpoint_occs) <= 1:
                continue
            upstream_rows = [m['row'] for m in endpoint_occs if m['is_tail']]
            downstream_rows = [m['row'] for m in endpoint_occs if m['is_head']]
            self.junctions[sec_id] = {
                'node_sec_id': sec_id,
                'rows': [m['row'] for m in endpoint_occs],
                'endpoint_meta': endpoint_occs,
                'upstream_rows': upstream_rows,
                'downstream_rows': downstream_rows,
            }

        for node_id, info in self.junctions.items():
            print(
                f"   Junction {node_id}: upstream="
                f"{len(info['upstream_rows'])}, downstream="
                f"{len(info['downstream_rows'])}, endpoints={len(info['rows'])}",
                flush=True)

        unresolved = []
        for sec_id, occs in self.section_occurrences.items():
            if len(occs) > 1:
                endpoint_cnt = sum(1 for m in occs if m['is_head'] or m['is_tail'])
                if endpoint_cnt <= 1:
                    unresolved.append(sec_id)
        if unresolved:
            print(
                "   Repeated sections without endpoint connections; check "
                f"the network table: {unresolved}", flush=True)

    def _build_static_arrays(self):
        N = len(self.row_sec_ids)
        self.section_to_idx = dict(self.section_representative_row)
        self.section_z_bed = {
            int(sec): float(self.dl.cross_sections[int(sec)]['Z_min'])
            for sec in self.section_representative_row.keys()
        }
        self.section_map = {
            int(sec): float(self.row_global_x[row])
            for sec, row in self.section_representative_row.items()
        }

        self.x_coords = np.asarray(self.row_global_x, dtype=float)
        self.x_local_array = np.asarray(self.row_local_x, dtype=float)
        self.x_global_array = np.asarray(self.row_global_x, dtype=float)
        self.reach_id_array = np.asarray(self.row_reach_ids, dtype=float)
        self.z_bed_array = np.asarray(self.row_z_bed, dtype=float)
        self.manning_n_array = np.asarray(self.row_manning_n, dtype=float)
        self.S0_array = np.zeros(N, dtype=float)
        self.reach_dx_tensor.clear()

        total_len = 0.0
        for reach_name in self.reach_names:
            rows = self.reach_row_indices[reach_name]
            x = self.reach_local_x[reach_name]
            zb = self.z_bed_array[rows]
            if len(rows) < 2:
                continue
            dx = np.diff(x)
            if np.any(dx <= 0):
                raise ValueError(
                    f"Section distances in reach {reach_name} must increase")
            dz = np.diff(zb)
            s_cell = -dz / (dx + 1e-8)
            local_s = np.zeros(len(rows), dtype=float)
            local_s[0] = s_cell[0]
            local_s[-1] = s_cell[-1]
            if len(rows) > 2:
                local_s[1:-1] = 0.5 * (s_cell[:-1] + s_cell[1:])
            self.S0_array[rows] = local_s
            self.reach_dx_tensor[reach_name] = torch.tensor(dx, dtype=torch.float32, device=self.device)
            total_len += float(np.sum(dx))

        self.L = max(float(total_len), 1.0)
        self.z_bed_tensor = torch.tensor(self.z_bed_array, dtype=torch.float32, device=self.device)
        self.sec_idx_tensor = torch.arange(N, device=self.device)
        print(
            f"Static network: {len(self.reach_names)} reaches, {N} nodes, "
            f"total reach length about {self.L:.1f} m", flush=True)

    @staticmethod
    def _multi_source_dijkstra(adjacency, sources):
        """Compute weighted multi-source shortest paths on the river graph."""
        n_nodes = len(adjacency)
        distance = np.full(n_nodes, np.inf, dtype=float)
        queue = []
        for source in sources:
            source = int(source)
            if 0 <= source < n_nodes and distance[source] > 0.0:
                distance[source] = 0.0
                heapq.heappush(queue, (0.0, source))

        while queue:
            dist_u, u = heapq.heappop(queue)
            if dist_u > distance[u]:
                continue
            for v, weight in adjacency[u]:
                candidate = dist_u + max(float(weight), 0.0)
                if candidate < distance[v]:
                    distance[v] = candidate
                    heapq.heappush(queue, (candidate, v))
        return distance

    def _build_physical_topology_features(self):
        """
        Build topology metrics from physical reach distances.

        Directed edges follow downstream flow. Zero-length junction edges join
        upstream reach tails to downstream reach heads.
        """
        n_nodes = len(self.row_sec_ids)
        downstream_graph = [[] for _ in range(n_nodes)]
        upstream_graph = [[] for _ in range(n_nodes)]
        undirected_graph = [[] for _ in range(n_nodes)]

        for reach_name in self.reach_names:
            rows = self.reach_row_indices[reach_name]
            x_local = self.reach_local_x[reach_name]
            for i in range(len(rows) - 1):
                u, v = int(rows[i]), int(rows[i + 1])
                distance = float(x_local[i + 1] - x_local[i])
                downstream_graph[u].append((v, distance))
                upstream_graph[v].append((u, distance))
                undirected_graph[u].append((v, distance))
                undirected_graph[v].append((u, distance))

        junction_degree = np.zeros(n_nodes, dtype=float)
        junction_rows = []
        for node in self.junctions.values():
            rows = [int(r) for r in node['rows']]
            junction_rows.extend(rows)
            degree = float(len(rows))
            for row in rows:
                junction_degree[row] = degree
            for upstream_row in node['upstream_rows']:
                for downstream_row in node['downstream_rows']:
                    u, v = int(upstream_row), int(downstream_row)
                    downstream_graph[u].append((v, 0.0))
                    upstream_graph[v].append((u, 0.0))
            # Junction occurrences represent one physical node.
            for i, u in enumerate(rows):
                for v in rows[i + 1:]:
                    undirected_graph[u].append((v, 0.0))
                    undirected_graph[v].append((u, 0.0))

        upper_rows = []
        for sec_id in self.upper_boundary_sections_sorted:
            upper_rows.extend(self.row_indices_by_section.get(sec_id, []))
        lower_rows = []
        for sec_id in self.lower_boundary_sections_sorted:
            lower_rows.extend(self.row_indices_by_section.get(sec_id, []))

        self.distance_from_upstream_m = self._multi_source_dijkstra(
            downstream_graph, upper_rows)
        self.distance_to_downstream_m = self._multi_source_dijkstra(
            upstream_graph, lower_rows)
        self.distance_to_junction_m = self._multi_source_dijkstra(
            undirected_graph, junction_rows)

        # Upstream signals use forward paths; downstream stages use reverse paths.
        route_columns = []
        route_names = []
        upper_count = 0
        for sec_id in self.upper_boundary_sections_sorted:
            sources = self.row_indices_by_section.get(sec_id, [])
            route_columns.append(self._multi_source_dijkstra(
                downstream_graph, sources))
            route_names.append(f'up_{sec_id}_q')
            upper_count += 1
        for sec_id in self.lower_boundary_sections_sorted:
            sources = self.row_indices_by_section.get(sec_id, [])
            route_columns.append(self._multi_source_dijkstra(
                upstream_graph, sources))
            route_names.append(f'low_{sec_id}_h')

        route_weights = np.zeros(
            (n_nodes, len(route_columns)), dtype=np.float32)
        route_lag_fraction = np.zeros_like(route_weights)
        for j, distances in enumerate(route_columns):
            finite = np.isfinite(distances)
            positive = distances[finite & (distances > 0)]
            scale = (
                float(np.quantile(positive, 0.95))
                if positive.size > 0 else 1.0
            )
            route_lag_fraction[finite, j] = np.clip(
                distances[finite] / max(scale, 1.0), 0.0, 1.0)
        for group_start, group_end in [
                (0, upper_count), (upper_count, len(route_columns))]:
            if group_end <= group_start:
                continue
            group_distances = route_columns[group_start:group_end]
            raw = np.zeros(
                (n_nodes, group_end - group_start), dtype=float)
            for j, distances in enumerate(group_distances):
                finite = np.isfinite(distances)
                positive = distances[finite & (distances > 0)]
                scale = (
                    float(np.quantile(positive, 0.75))
                    if positive.size > 0 else 1.0
                )
                scale = max(scale, 1.0)
                raw[finite, j] = np.exp(-distances[finite] / scale)
            denominator = raw.sum(axis=1, keepdims=True)
            normalized = np.divide(
                raw, denominator,
                out=np.zeros_like(raw),
                where=denominator > 0,
            )
            route_weights[:, group_start:group_end] = normalized

        self.boundary_route_names = route_names
        self.boundary_route_weights_array = route_weights
        self.boundary_lag_fraction_array = route_lag_fraction

        # Fall back to undirected physical paths for incomplete direction data.
        undirected_up = self._multi_source_dijkstra(
            undirected_graph, upper_rows)
        undirected_low = self._multi_source_dijkstra(
            undirected_graph, lower_rows)
        self.distance_from_upstream_m = np.where(
            np.isfinite(self.distance_from_upstream_m),
            self.distance_from_upstream_m, undirected_up)
        self.distance_to_downstream_m = np.where(
            np.isfinite(self.distance_to_downstream_m),
            self.distance_to_downstream_m, undirected_low)

        if not junction_rows:
            self.distance_to_junction_m = np.zeros(n_nodes, dtype=float)
        else:
            finite = np.isfinite(self.distance_to_junction_m)
            fallback = (
                float(np.max(self.distance_to_junction_m[finite]))
                if np.any(finite) else 0.0
            )
            self.distance_to_junction_m = np.where(
                finite, self.distance_to_junction_m, fallback)

        self.junction_degree_array = junction_degree
        print(
            "DeepONet topology distances use physical reach lengths and "
            "junction connections",
            flush=True,
        )

    def _build_lookup_table_if_needed(self):
        if self.lut.built:
            return
        print("Building hydraulic lookup tables", flush=True)
        self.sec_id_to_lut_idx = {}
        unique_sections = list(self.section_representative_row.keys())
        for i, sec_id in enumerate(unique_sections):
            sec_id = int(sec_id)
            self.sec_id_to_lut_idx[sec_id] = i
            z_min = float(self.dl.cross_sections[sec_id]["Z_min"])
            # Extend lookup tables above surveyed banks using the selected mode.
            z_max = float(self.calc.geometry_upper_level(sec_id))
            if z_max <= z_min + 1e-3:
                z_max = z_min + 1.0
            ws = np.linspace(z_min, z_max, 81)
            As = [self.calc.compute_ABK(sec_id, w)[0] for w in ws]
            Bs = [self.calc.compute_ABK(sec_id, w)[1] for w in ws]
            Ks = [self.calc.compute_ABK(sec_id, w)[2] for w in ws]

            As = np.asarray(As, dtype=float)
            Bs = np.asarray(Bs, dtype=float)
            Ks = np.asarray(Ks, dtype=float)
            # Enforce nondecreasing area and conveyance.
            As = np.maximum.accumulate(As)
            Ks = np.maximum.accumulate(Ks)
            Bs = np.maximum(Bs, 1e-3)
            self.lut.add_section(float(i), {'water_levels': ws, 'A': As, 'B': Bs, 'K': Ks})
        self.row_lut_indices = [self.sec_id_to_lut_idx[int(sec)] for sec in self.row_sec_ids]
        self.row_lut_indices_tensor = torch.tensor(self.row_lut_indices, dtype=torch.long, device=self.device)
        self.lut.n_levels = 81
        self.lut.build()
        print(f"   Overbank mode={self.overbank_mode}, side_slope={self.overbank_side_slope}, "
              f"height={self.overbank_height} m", flush=True)

    # ----------------------------------------------------------
    # File parsing
    # ----------------------------------------------------------

    @staticmethod
    def _to_rel_hours(times):
        if pd.api.types.is_datetime64_any_dtype(times):
            t0 = times.min()
            return ((times - t0).dt.total_seconds() / 3600.0).values.astype(float)
        return pd.to_numeric(times, errors='coerce').fillna(0.0).values.astype(float)

    @staticmethod
    def _csv_path(file_path):
        """Return a validated CSV path, preserving optional None values."""
        if file_path is None:
            return None
        path = str(file_path)
        if not path.lower().endswith(".csv"):
            raise ValueError(f"Only CSV input is supported: {path}")
        if not os.path.isfile(path):
            raise FileNotFoundError(path)
        return path

    @staticmethod
    def _pick_csv_column(columns, keywords, fallback_index=None):
        for column in columns:
            text = str(column).lower()
            if any(str(keyword).lower() in text for keyword in keywords):
                return column
        if fallback_index is not None and 0 <= fallback_index < len(columns):
            return columns[fallback_index]
        return None

    def _row_for_occurrence_name(self, name):
        """Resolve names such as 55_sec2_0 to unique network rows."""
        text = str(name).strip()
        parts = text.split('_')
        if len(parts) < 3:
            return None, None
        try:
            sec_id = int(float(parts[0]))
            local_idx = int(float(parts[-1]))
        except ValueError:
            return None, None
        reach_name = '_'.join(parts[1:-1])
        for meta in self.row_to_meta:
            if (int(meta['sec_id']) == sec_id
                    and str(meta['reach_name']) == reach_name
                    and int(meta['local_idx']) == local_idx):
                return int(meta['row']), sec_id
        return None, sec_id

    def _rows_for_node_identifier(self, identifier):
        """
        Resolve a CSV section identifier to network rows.

        A numeric section selects its representative row. A composite identifier
        such as 55_sec2_0 selects one explicit reach occurrence.
        """
        text = str(identifier).strip()
        if not text:
            return [], None

        occurrence_row, occurrence_sec = self._row_for_occurrence_name(text)
        if occurrence_row is not None:
            return [int(occurrence_row)], int(occurrence_sec)

        try:
            sec_id = int(float(text))
        except (TypeError, ValueError):
            return [], None

        representative = self.section_representative_row.get(sec_id)
        if representative is None:
            return [], sec_id
        return [int(representative)], sec_id

    def _filter_resolved_rows(self, rows, observed_filter, sec_id):
        """Filter resolved rows using the supervision configuration."""
        if observed_filter is None:
            return list(rows)
        representative = self.section_representative_row.get(int(sec_id))
        selected = []
        for row in rows:
            if row in observed_filter['rows']:
                selected.append(int(row))
            elif (int(sec_id) in observed_filter['sections']
                  and representative is not None
                  and int(row) == int(representative)):
                selected.append(int(row))
        return selected

    def _compile_observed_sections_filter(self, observed_sections):
        """Support numeric sections and composite junction-node names."""
        if observed_sections is None:
            return None
        section_ids, row_ids, unresolved = set(), set(), []
        for token in observed_sections:
            text = str(token).strip()
            if not text:
                continue
            row, _ = self._row_for_occurrence_name(text)
            if row is not None:
                row_ids.add(row)
                continue
            try:
                section_ids.add(int(float(text)))
            except ValueError:
                unresolved.append(text)
        if unresolved:
            available = [
                self._sheet_name_for_row(i)
                for i in range(len(self.row_sec_ids))
                if len(self.row_indices_by_section.get(self.row_sec_ids[i], [])) > 1
            ]
            raise ValueError(
                f"Unrecognized observation or junction nodes: {unresolved}. "
                f"Examples of repeated nodes: {available[:20]}"
            )
        return {'sections': section_ids, 'rows': row_ids}

    def _observed_rows_for_section(self, observed_filter, sec_id):
        rows = list(self.row_indices_by_section.get(int(sec_id), []))
        if observed_filter is None:
            return rows
        selected = [r for r in rows if r in observed_filter['rows']]
        if int(sec_id) in observed_filter['sections']:
            # Composite names select a specific repeated junction occurrence.
            rep = self.section_representative_row.get(int(sec_id))
            if rep is not None and int(rep) not in selected:
                selected.insert(0, int(rep))
        return selected

    def _parse_boundary_file(self, file_path):
        """
        Read a boundary CSV. The model uses upstream discharge and downstream
        water level as boundary inputs.
        """
        upper_sections = set(self.upper_boundary_sections_sorted)
        lower_sections = set(self.lower_boundary_sections_sorted)
        csv_path = self._csv_path(file_path)
        df = pd.read_csv(csv_path)
        if df.empty:
            raise ValueError(f"Boundary CSV is empty: {csv_path}")

        cols = list(df.columns)
        sec_col = self._pick_csv_column(
            cols, ['section_id'], 0)
        type_col = self._pick_csv_column(
            cols, ['boundary_type'])
        time_col = self._pick_csv_column(
            cols, ['time_hours', 'time'],
            2 if type_col is not None else 1)
        q_col = self._pick_csv_column(
            cols, ['discharge_m3s', 'discharge'],
            3 if type_col is not None else 2)
        z_col = self._pick_csv_column(
            cols, ['water_level_m', 'water_level'],
            4 if type_col is not None else 3)

        upper_data, lower_data = {}, {}
        tmp = df.copy()
        tmp[sec_col] = pd.to_numeric(
            tmp[sec_col], errors='coerce').astype('Int64')
        tmp = tmp.dropna(subset=[sec_col])
        for sid, group in tmp.groupby(sec_col, sort=False):
            sec_id = int(sid)
            t = self._to_rel_hours(group[time_col])
            q = (
                pd.to_numeric(group[q_col], errors='coerce')
                .ffill().bfill().fillna(0.0).values.astype(float)
                if q_col is not None else np.zeros_like(t)
            )
            z = (
                pd.to_numeric(group[z_col], errors='coerce')
                .ffill().bfill().fillna(0.0).values.astype(float)
                if z_col is not None else np.zeros_like(t)
            )
            boundary_text = (
                ' '.join(
                    str(v).lower()
                    for v in group[type_col].dropna().unique())
                if type_col is not None else ''
            )
            is_downstream = (
                'downstream' in boundary_text)
            is_upstream = (
                'upstream' in boundary_text)
            if sec_id in upper_sections and not is_downstream:
                upper_data[sec_id] = {'time': t, 'discharge': q, 'water_level': z}
            if sec_id in lower_sections and not is_upstream:
                lower_data[sec_id] = {'time': t, 'discharge': q, 'water_level': z}

        missing_upper = sorted(upper_sections - set(upper_data))
        missing_lower = sorted(lower_sections - set(lower_data))
        if missing_upper or missing_lower:
            raise ValueError(
                f"Boundary CSV {csv_path} is incomplete: "
                f"missing upstream={missing_upper}, downstream={missing_lower}"
            )
        return upper_data, lower_data

    def _parse_lateral_file(self, file_path):
        if file_path is None:
            return {}
        csv_path = self._csv_path(file_path)
        df = pd.read_csv(csv_path)
        if df.empty:
            return {}

        cols = list(df.columns)
        sec_col = self._pick_csv_column(
            cols, ['section_id'], 0)
        time_col = self._pick_csv_column(
            cols, ['time_hours', 'time'], 1)
        q_col = self._pick_csv_column(
            cols, ['discharge_m3s', 'discharge'], 2)
        lateral = {}
        tmp = df[[sec_col, time_col, q_col]].copy()
        tmp.columns = ['section_id', 'time_hours', 'discharge_m3s']
        tmp['section_id'] = pd.to_numeric(
            tmp['section_id'], errors='coerce').astype('Int64')
        for sec_id, group in tmp.dropna(
                subset=['section_id']).groupby('section_id', sort=False):
            lateral[int(sec_id)] = {
                'time': self._to_rel_hours(group['time_hours']),
                'discharge': (
                    pd.to_numeric(group['discharge_m3s'], errors='coerce')
                    .ffill().bfill().fillna(0.0).values.astype(float)
                ),
            }
        return lateral

    def _parse_result_file_sparse(self, file_path, observed_sections_set=None):
        if file_path is None:
            return None, None, None, []
        csv_path = self._csv_path(file_path)
        df = pd.read_csv(csv_path)
        if df.empty:
            return None, None, None, []

        cols = list(df.columns)
        sec_col = self._pick_csv_column(
            cols, ['section_id'], 0)
        time_col = self._pick_csv_column(
            cols, ['time_hours', 'time'], 1)
        z_col = self._pick_csv_column(
            cols, ['water_level_m', 'water_level'], 2)
        q_col = self._pick_csv_column(
            cols, ['discharge_m3s', 'discharge'], 3)
        row_series, t_union = {}, []
        unresolved_identifiers = []
        for sid, group in df.dropna(
                subset=[sec_col]).groupby(sec_col, sort=False):
            resolved_rows, sec_id = self._rows_for_node_identifier(sid)
            if not resolved_rows or sec_id is None:
                unresolved_identifiers.append(str(sid))
                continue
            target_rows = self._filter_resolved_rows(
                resolved_rows, observed_sections_set, sec_id)
            if not target_rows:
                continue
            t = self._to_rel_hours(group[time_col])
            z = (
                pd.to_numeric(group[z_col], errors='coerce')
                .ffill().bfill().values.astype(float)
            )
            q = (
                pd.to_numeric(group[q_col], errors='coerce')
                .ffill().bfill().values.astype(float)
            )
            for target_row in target_rows:
                row_series[int(target_row)] = {
                    'time': t, 'z': z, 'q': q}
            t_union.extend(t.tolist())

        if unresolved_identifiers:
            preview = unresolved_identifiers[:10]
            print(
                f"Skipped {len(unresolved_identifiers)} unmapped target "
                f"identifiers: {preview}",
                flush=True,
            )

        if not row_series:
            return None, None, None, []
        t_union = np.array(
            sorted(set(float(v) for v in t_union)), dtype=float)
        N = len(self.row_sec_ids)
        z_grid = np.zeros((N, len(t_union)), dtype=float)
        q_grid = np.zeros((N, len(t_union)), dtype=float)
        obs_row_indices = []
        for row in range(N):
            if row in row_series:
                item = row_series[row]
                z_grid[row, :] = np.interp(t_union, item['time'], item['z'])
                q_grid[row, :] = np.interp(t_union, item['time'], item['q'])
                obs_row_indices.append(row)
        return t_union, z_grid, q_grid, obs_row_indices

    def _parse_initial_condition_file(self, file_path):
        if file_path is None:
            return None, None
        csv_path = self._csv_path(file_path)
        df = pd.read_csv(csv_path)
        df.columns = [
            'section_id', 'initial_discharge_m3s', 'initial_water_level_m']
        z0 = np.full(len(self.row_sec_ids), np.nan)
        q0 = np.full(len(self.row_sec_ids), np.nan)
        unresolved_identifiers = []
        for _, row in df.dropna(subset=['section_id']).iterrows():
            resolved_rows, _ = self._rows_for_node_identifier(
                row['section_id'])
            if not resolved_rows:
                unresolved_identifiers.append(str(row['section_id']))
                continue
            q_value = pd.to_numeric(
                pd.Series([row['initial_discharge_m3s']]),
                errors='coerce').iloc[0]
            z_value = pd.to_numeric(
                pd.Series([row['initial_water_level_m']]),
                errors='coerce').iloc[0]
            if pd.isna(q_value) or pd.isna(z_value):
                continue
            for k in resolved_rows:
                q0[k] = float(q_value)
                z0[k] = float(z_value)

        if unresolved_identifiers:
            preview = unresolved_identifiers[:10]
            print(
                f"Skipped {len(unresolved_identifiers)} unmapped initial-"
                f"condition identifiers: {preview}",
                flush=True,
            )
        valid = ~np.isnan(z0)
        if valid.sum() == 0:
            return None, None
        idx_v = np.where(valid)[0]
        idx_m = np.where(~valid)[0]
        if valid.sum() >= 2 and len(idx_m) > 0:
            z0[idx_m] = np.interp(idx_m, idx_v, z0[idx_v])
            q0[idx_m] = np.interp(idx_m, idx_v, q0[idx_v])
        elif valid.sum() == 1 and len(idx_m) > 0:
            z0[idx_m] = z0[idx_v[0]]
            q0[idx_m] = q0[idx_v[0]]
        return z0, q0

    def _load_single_event(self, spec):
        upper_data_raw, lower_data_raw = self._parse_boundary_file(spec.boundary_file)
        lateral_dict = self._parse_lateral_file(spec.lateral_inflow_file)

        # ------------------------------------------------------
        # Keep supervision targets separate from evaluation-only references.
        # ------------------------------------------------------
        # Reference targets are loaded for evaluation and result export only.
        t_ref, z_ref, q_ref, ref_indices = self._parse_result_file_sparse(
            spec.result_file, observed_sections_set=None)

        # None means full supervision, [] means no observation supervision,
        # and a list selects sparse supervised sections.
        if spec.observed_sections is None:
            t_res, z_target, q_target, obs_indices = t_ref, z_ref, q_ref, list(ref_indices or [])
        elif len(spec.observed_sections) == 0:
            t_res, z_target, q_target, obs_indices = None, None, None, []
        else:
            obs_set = self._compile_observed_sections_filter(
                spec.observed_sections)
            t_res, z_target, q_target, obs_indices = self._parse_result_file_sparse(
                spec.result_file, observed_sections_set=obs_set)

        # Reference times are included only to align later evaluation.
        t_union = []
        if t_ref is not None:
            t_union.extend(t_ref.tolist())
        if t_res is not None:
            t_union.extend(t_res.tolist())
        for item in list(upper_data_raw.values()) + list(lower_data_raw.values()):
            t_union.extend(item['time'].tolist())
        for item in lateral_dict.values():
            t_union.extend(item['time'].tolist())
        t_union = np.array(sorted(set([float(v) for v in t_union])), dtype=float)
        if len(t_union) < 2:
            raise ValueError(f"Event {spec.name} has too few time steps")

        upper_q, lower_z = {}, {}
        for sec, item in upper_data_raw.items():
            upper_q[sec] = np.interp(t_union, item['time'], item['discharge'])
        for sec, item in lower_data_raw.items():
            lower_z[sec] = np.interp(t_union, item['time'], item['water_level'])

        first_up = self.upper_boundary_sections_sorted[0]
        first_low = self.lower_boundary_sections_sorted[0]
        q_up = upper_q[first_up]
        z_low = lower_z[first_low]

        N = len(self.row_sec_ids)
        lateral = np.zeros((N, len(t_union)), dtype=float)
        node_source = np.zeros((N, len(t_union)), dtype=float)
        for sec_id, item in lateral_dict.items():
            if sec_id not in self.row_indices_by_section:
                print(
                    f"   Lateral-inflow section {sec_id} is not in the "
                    "network and was skipped", flush=True)
                continue
            q_lat = np.interp(t_union, item['time'], item['discharge'])
            if sec_id in self.junctions:
                for r in self.junctions[sec_id]['rows']:
                    node_source[r, :] += q_lat
                print(f"   Section {sec_id}: lateral source applied", flush=True)
            else:
                for r in self.row_indices_by_section[sec_id]:
                    lateral[r, :] += q_lat

        # Interpolate supervision targets onto the shared event time axis.
        if z_target is not None and q_target is not None:
            z_t = np.zeros((N, len(t_union)), dtype=float)
            q_t = np.zeros((N, len(t_union)), dtype=float)
            for i in range(N):
                if i in obs_indices and t_res is not None:
                    z_t[i, :] = np.interp(t_union, t_res, z_target[i])
                    q_t[i, :] = np.interp(t_union, t_res, q_target[i])
            z_target, q_target = z_t, q_t

        # Interpolate full references for evaluation only.
        reference_z, reference_q = None, None
        reference_indices = list(ref_indices or [])
        if z_ref is not None and q_ref is not None:
            z_r = np.zeros((N, len(t_union)), dtype=float)
            q_r = np.zeros((N, len(t_union)), dtype=float)
            for i in range(N):
                if i in reference_indices and t_ref is not None:
                    z_r[i, :] = np.interp(t_union, t_ref, z_ref[i])
                    q_r[i, :] = np.interp(t_union, t_ref, q_ref[i])
            reference_z, reference_q = z_r, q_r

        z0, q0 = self._parse_initial_condition_file(spec.initial_condition_file)
        if z0 is None or q0 is None:
            if z_target is not None and len(obs_indices) > 0:
                if len(obs_indices) >= 2:
                    z0 = np.interp(np.arange(N), obs_indices, z_target[obs_indices, 0])
                else:
                    z0 = self.z_bed_array + 1.0
                q0 = np.repeat(q_up[0], N)
            elif reference_z is not None and len(reference_indices) > 0:
                # A reference at t=0 may initialize an otherwise missing state.
                z0 = reference_z[:, 0].copy()
                q0 = reference_q[:, 0].copy()
            else:
                z0 = self.z_bed_array + 1.0
                q0 = np.repeat(q_up[0], N)

        n_obs = len(set(obs_indices))
        n_total = N
        n_ref = len(set(reference_indices))
        if spec.observed_sections is not None and len(spec.observed_sections) == 0:
            obs_info = (
                f"no observations (reference={n_ref}/{n_total})"
                if spec.result_file else "no observations")
        else:
            obs_info = (
                f"observed nodes={n_obs}/{n_total}"
                if n_obs < n_total else "fully supervised")
            if n_obs == 0:
                obs_info = (
                    f"no observations (reference={n_ref}/{n_total})"
                    if n_ref > 0 else "no observations")

        evt = RawEvent(
            name=spec.name, t_hours=t_union,
            upper_q=upper_q, lower_z=lower_z,
            lateral=lateral, node_source=node_source,
            z0=z0, q0=q0,
            target_z=z_target, target_q=q_target,
            reference_z=reference_z, reference_q=reference_q,
            observed_section_indices=obs_indices,
            reference_section_indices=reference_indices)
        print(f"{evt.name}: T={len(t_union)}, {obs_info}, "
              f"Q_up({first_up})=[{q_up.min():.1f},{q_up.max():.1f}], "
              f"Z_low({first_low})=[{z_low.min():.2f},{z_low.max():.2f}]", flush=True)
        return evt

    def _load_all_events(self):
        self.raw_events = [self._load_single_event(spec) for spec in self.event_specs]

    # ----------------------------------------------------------
    # Normalization and input construction
    # ----------------------------------------------------------

    def _event_scale_features(self, evt: RawEvent):
        """
        Compute event-scale features from boundaries, initial states, and sources.
        Targets are excluded to prevent information leakage.
        """
        q_candidates = []
        # Only upstream discharge is available during prediction.
        for arr in evt.upper_q.values():
            q_candidates.append(np.asarray(arr, dtype=float).reshape(-1))
        if evt.q0 is not None and len(evt.q0) > 0:
            q_candidates.append(np.asarray(evt.q0, dtype=float).reshape(-1))
        if evt.lateral.size > 0:
            vals = evt.lateral[np.abs(evt.lateral) > 1e-8]
            if vals.size > 0:
                q_candidates.append(vals.reshape(-1))
        if evt.node_source.size > 0:
            vals = evt.node_source[np.abs(evt.node_source) > 1e-8]
            if vals.size > 0:
                q_candidates.append(vals.reshape(-1))

        if q_candidates:
            q_all = np.concatenate(q_candidates)
            q_all = q_all[np.isfinite(q_all)]
            q_abs_max = float(np.max(np.abs(q_all))) if q_all.size > 0 else 1.0
        else:
            q_abs_max = 1.0
        q_abs_max = max(q_abs_max, 1.0)
        q_scale_evt = max(q_abs_max, 20.0)

        h_range_evt = 0.1
        # Only downstream stage/depth is available during prediction.
        for sec, z in evt.lower_z.items():
            if sec not in self.section_z_bed:
                continue
            bed = self.section_z_bed[sec]
            h = np.asarray(z, dtype=float) - bed
            if np.all(np.isfinite(h)):
                h_range_evt = max(h_range_evt, float(h.max() - h.min()))
        if evt.z0 is not None and len(evt.z0) > 0:
            h0 = np.asarray(evt.z0, dtype=float) - self.z_bed_array
            h0 = h0[np.isfinite(h0)]
            if h0.size > 0:
                h_range_evt = max(h_range_evt, float(h0.max() - h0.min()) * 0.25)
        h_range_evt = max(h_range_evt, 0.1)

        return {
            'q_abs_max': q_abs_max,
            'q_scale_evt': q_scale_evt,
            'h_range_evt': h_range_evt,
            'log_qmax': np.log10(q_abs_max),
            'log_hrange': np.log10(h_range_evt),
        }

    def _compute_normalization(self):
        """
        Estimate normalization parameters from training events only.

        State variables use affine scaling based on the training conditions:
          q_norm = (Q - q_center) / q_scale
          h_norm = (h - h_center) / h_scale
          source_norm = (q_lat - source_center) / source_scale
        Physical evaluation metrics use de-normalized values.
        """
        if self.train_event_names:
            train_name_set = set(self.train_event_names)
            norm_events = [evt for evt in self.raw_events if evt.name in train_name_set]
            if len(norm_events) == 0:
                print(
                    "Training event names did not match loaded events; "
                    "normalization uses all events", flush=True)
                norm_events = self.raw_events
        else:
            norm_events = self.raw_events

        self._evt_log_qmax = []
        self._evt_log_hrange = []
        self._evt_q_scale = []
        self._evt_h_range = []
        for evt in self.raw_events:
            feat = self._event_scale_features(evt)
            self._evt_log_qmax.append(feat['log_qmax'])
            self._evt_log_hrange.append(feat['log_hrange'])
            self._evt_q_scale.append(feat['q_scale_evt'])
            self._evt_h_range.append(feat['h_range_evt'])

        all_q, all_h, all_source, all_duration = [], [], [], []
        for evt in norm_events:
            for arr in evt.upper_q.values():
                all_q.extend(np.asarray(arr, dtype=float).reshape(-1).tolist())
            all_q.extend(np.asarray(evt.q0, dtype=float).reshape(-1).tolist())

            if evt.lateral.size > 0:
                vals = evt.lateral[np.isfinite(evt.lateral)]
                all_source.extend(vals.reshape(-1).tolist())
                nz = vals[np.abs(vals) > 1e-8]
                if nz.size > 0:
                    all_q.extend(nz.tolist())
            if evt.node_source.size > 0:
                vals = evt.node_source[np.isfinite(evt.node_source)]
                all_source.extend(vals.reshape(-1).tolist())
                nz = vals[np.abs(vals) > 1e-8]
                if nz.size > 0:
                    all_q.extend(nz.tolist())

            if evt.target_q is not None:
                if evt.observed_section_indices is None:
                    all_q.extend(evt.target_q.reshape(-1).tolist())
                elif len(evt.observed_section_indices) > 0:
                    for idx in evt.observed_section_indices:
                        all_q.extend(evt.target_q[idx].tolist())

            for sec, arr in evt.lower_z.items():
                bed = self.section_z_bed.get(sec, self.z_bed_array[self.section_representative_row[sec]])
                all_h.extend((np.asarray(arr, dtype=float) - bed).tolist())
            all_h.extend((evt.z0 - self.z_bed_array).tolist())
            if evt.target_z is not None:
                if evt.observed_section_indices is None:
                    all_h.extend((evt.target_z - self.z_bed_array[:, None]).reshape(-1).tolist())
                elif len(evt.observed_section_indices) > 0:
                    for idx in evt.observed_section_indices:
                        all_h.extend((evt.target_z[idx] - self.z_bed_array[idx]).tolist())

            duration = max(float(evt.t_hours[-1] - evt.t_hours[0]), 1e-6)
            all_duration.append(duration)

        valid_q = np.asarray(all_q, dtype=float)
        valid_q = valid_q[np.isfinite(valid_q)]
        valid_h = np.asarray(all_h, dtype=float)
        valid_h = valid_h[np.isfinite(valid_h)]
        valid_source = np.asarray(all_source, dtype=float)
        valid_source = valid_source[np.isfinite(valid_source)]
        all_duration = np.asarray(all_duration, dtype=float)
        all_duration = all_duration[np.isfinite(all_duration)]

        def _affine_range(vals, default_min, default_max, min_half, margin_ratio=0.05):
            vals = np.asarray(vals, dtype=float)
            vals = vals[np.isfinite(vals)]
            if vals.size == 0:
                lo, hi = float(default_min), float(default_max)
            else:
                lo, hi = float(np.min(vals)), float(np.max(vals))
                span = max(hi - lo, min_half * 2.0)
                margin = max(span * margin_ratio, min_half * 0.05)
                lo -= margin
                hi += margin
            if hi <= lo + 1e-12:
                hi = lo + 2.0 * min_half
            center = 0.5 * (lo + hi)
            scale = max(0.5 * (hi - lo), min_half)
            return lo, hi, center, scale

        self.q_min, self.q_max, self.q_center, self.q_scale = _affine_range(
            valid_q, 0.0, 1.0, min_half=1.0, margin_ratio=0.05)
        self.h_min, self.h_max, self.h_center, self.h_scale = _affine_range(
            valid_h, 0.05, 2.0, min_half=0.25, margin_ratio=0.05)
        self.source_min, self.source_max, self.source_center, self.source_scale = _affine_range(
            valid_source, 0.0, 1.0, min_half=0.1, margin_ratio=0.05)
        self.duration_scale = max(float(np.quantile(all_duration, 0.95)), 1.0) if all_duration.size > 0 else 1.0

        xg_min, xg_max = self.x_global_array.min(), self.x_global_array.max()
        self.x_global_norm = (self.x_global_array - xg_min) / max(xg_max - xg_min, 1.0) * 2 - 1
        self.x_norm = self.x_global_norm
        self.x_local_norm = np.zeros_like(self.x_global_norm)
        for reach_name in self.reach_names:
            rows = self.reach_row_indices[reach_name]
            xl = self.x_local_array[rows]
            self.x_local_norm[rows] = xl / max(float(xl.max() - xl.min()), 1.0) * 2 - 1
        self.reach_norm = self.reach_id_array / max(float(max(len(self.reach_names) - 1, 1)), 1.0) * 2 - 1
        self.z_bed_norm = (self.z_bed_array - self.z_bed_array.min()) / max(self.z_bed_array.max() - self.z_bed_array.min(), 1.0)
        s0_scale = max(np.abs(self.S0_array).max(), 1e-6)
        self.S0_norm = self.S0_array / s0_scale
        n_scale = max(self.manning_n_array.max(), 0.01)
        self.n_norm = self.manning_n_array / n_scale

        raw_name_to_idx = {evt.name: i for i, evt in enumerate(self.raw_events)}
        norm_indices = [raw_name_to_idx[evt.name] for evt in norm_events if evt.name in raw_name_to_idx]
        if len(norm_indices) == 0:
            norm_indices = list(range(len(self.raw_events)))
        all_lq = np.asarray([self._evt_log_qmax[i] for i in norm_indices], dtype=float)
        all_lh = np.asarray([self._evt_log_hrange[i] for i in norm_indices], dtype=float)
        self.log_qmax_center = float(np.mean(all_lq)) if all_lq.size > 0 else 0.0
        self.log_qmax_scale = max(float(np.std(all_lq)), 0.5) if all_lq.size > 0 else 1.0
        self.log_hrange_center = float(np.mean(all_lh)) if all_lh.size > 0 else 0.0
        self.log_hrange_scale = max(float(np.std(all_lh)), 0.3) if all_lh.size > 0 else 1.0

        all_lq_all = np.asarray(self._evt_log_qmax, dtype=float)
        all_lh_all = np.asarray(self._evt_log_hrange, dtype=float)
        print(f"Training-set normalization: Q=[{self.q_min:.2f},{self.q_max:.2f}], "
              f"h=[{self.h_min:.2f},{self.h_max:.2f}], "
              f"source=[{self.source_min:.3g},{self.source_max:.3g}], duration_scale={self.duration_scale:.2f} h", flush=True)
        print(f"   Center/scale: q_center={self.q_center:.2f}, q_scale={self.q_scale:.2f}, "
              f"h_center={self.h_center:.2f}, h_scale={self.h_scale:.2f}, "
              f"source_center={self.source_center:.3g}, source_scale={self.source_scale:.3g}", flush=True)
        print(f"   conditioning(all events): log_qmax=[{all_lq_all.min():.2f},{all_lq_all.max():.2f}], "
              f"log_hrange=[{all_lh_all.min():.2f},{all_lh_all.max():.2f}]", flush=True)

    def _norm_q(self, q):
        return (q - self.q_center) / self.q_scale

    def _denorm_q(self, qn):
        return qn * self.q_scale + self.q_center

    def _norm_q_evt(self, q, q_scale_evt):
        q_scale_evt = max(float(q_scale_evt), 1e-6)
        return np.asarray(q, dtype=float) / q_scale_evt

    def _norm_source(self, q):
        return (q - self.source_center) / self.source_scale

    def _norm_source_evt(self, q, q_scale_evt):
        q_scale_evt = max(float(q_scale_evt), 1e-6)
        return np.asarray(q, dtype=float) / q_scale_evt

    def _denorm_source(self, qn):
        return qn * self.source_scale + self.source_center

    def _norm_h(self, h):
        return (h - self.h_center) / self.h_scale

    def _denorm_h(self, hn):
        return hn * self.h_scale + self.h_center

    def _norm_h_rel_evt(self, h, h_range_evt):
        h_range_evt = max(float(h_range_evt), 1e-6)
        h = np.asarray(h, dtype=float)
        return (h - h[0]) / h_range_evt

    def _resample_1d(self, t_old, y_old, tau_new, duration_h):
        tau_old = (t_old - t_old[0]) / max(duration_h, 1e-8)
        return np.interp(tau_new, tau_old, y_old)

    def _resample_2d(self, t_old, y_old, tau_new, duration_h):
        tau_old = (t_old - t_old[0]) / max(duration_h, 1e-8)
        out = np.zeros((y_old.shape[0], len(tau_new)), dtype=float)
        for i in range(y_old.shape[0]):
            out[i, :] = np.interp(tau_new, tau_old, y_old[i, :])
        return out

    def _set_time_grid_from_events(self):
        """
        Preserve native event time steps while using a shared tensor size.
        Shorter events are padded with their final value, and time_mask excludes
        padding from supervision and consistency regularization.
        """
        if len(self.raw_events) == 0:
            raise RuntimeError("Load events before configuring the time grid")
        max_T = max(len(evt.t_hours) for evt in self.raw_events)
        if self.n_time_model is None:
            self.n_time_model = int(max_T)
        elif self.n_time_model < max_T:
            print(
                f"n_time_model={self.n_time_model} is shorter than the "
                f"longest event; using {max_T}", flush=True)
            self.n_time_model = int(max_T)
        self.time_coord = np.linspace(0.0, 1.0, self.n_time_model, dtype=float)
        print(f"Time grid: padded event series, max_T={self.n_time_model}", flush=True)

    @staticmethod
    def _event_dt_hours(evt: RawEvent):
        t = np.asarray(evt.t_hours, dtype=float)
        if len(t) >= 2:
            dt = float(np.median(np.diff(t)))
            if np.isfinite(dt) and dt > 0:
                return dt
        return 1.0

    @staticmethod
    def _pad_1d_last(arr, Tm, fill_value=0.0):
        arr = np.asarray(arr, dtype=float).reshape(-1)
        out = np.zeros(Tm, dtype=float)
        L = min(len(arr), Tm)
        if L > 0:
            out[:L] = arr[:L]
            out[L:] = arr[L - 1]
        else:
            out[:] = fill_value
        return out

    @staticmethod
    def _pad_2d_last(arr, Tm, fill_value=0.0):
        arr = np.asarray(arr, dtype=float)
        if arr.ndim != 2:
            raise ValueError("_pad_2d_last expects a two-dimensional [N,T] array")
        N = arr.shape[0]
        out = np.zeros((N, Tm), dtype=float)
        L = min(arr.shape[1], Tm)
        if L > 0:
            out[:, :L] = arr[:, :L]
            out[:, L:] = arr[:, L - 1:L]
        else:
            out[:, :] = fill_value
        return out

    def _build_lateral_features(self, lateral_q):
        """Build local and cumulative lateral-source channels for a reach."""
        N, T = lateral_q.shape
        lat_density = np.zeros_like(lateral_q)
        lat_cum = np.zeros_like(lateral_q)

        for reach_name in self.reach_names:
            rows = self.reach_row_indices[reach_name]
            if len(rows) < 2:
                continue
            dx = self.reach_dx_tensor[reach_name].detach().cpu().numpy().reshape(-1)
            reach_lat = np.zeros((len(rows) - 1, T), dtype=float)
            reach_cum = np.zeros((len(rows), T), dtype=float)

            for j, row in enumerate(rows):
                q = lateral_q[row]
                if np.max(np.abs(q)) <= 1e-12:
                    continue
                cell = j if j < len(rows) - 1 else len(rows) - 2
                density = q / max(dx[cell], 1e-6)
                reach_lat[cell, :] += density
                reach_cum[j:, :] += q

            # Apply each cell source to its two neighboring section rows.
            for c in range(len(rows) - 1):
                lat_density[rows[c], :] += reach_lat[c, :]
                lat_density[rows[c + 1], :] += reach_lat[c, :]
            for j, row in enumerate(rows):
                lat_cum[row, :] += reach_cum[j, :]

        return lat_density, lat_cum

    def _build_input_grid_for_event(self, evt: RawEvent, i_evt_for_conditioning: Optional[int] = None):
        """
        Build the input grid for one event.

        Boundary inputs match prediction-time availability: upstream discharge
        and downstream stage/depth. Other channels include static topology,
        normalized time, masks, lateral sources, initial states, duration, and
        event-scale conditioning features.
        """
        if self.n_time_model is None:
            raise RuntimeError(
                "n_time_model is not set; call _set_time_grid_from_events()")
        Tm = self.n_time_model
        N = len(self.row_sec_ids)
        valid_T = min(len(evt.t_hours), Tm)
        dt_h = self._event_dt_hours(evt)
        duration_h = max(float(evt.t_hours[-1] - evt.t_hours[0]), dt_h)

        time_mask_1d = np.zeros(Tm, dtype=np.float32)
        time_mask_1d[:valid_T] = 1.0
        tau_model = np.linspace(0.0, 1.0, Tm, dtype=float)

        # Pad prediction-time boundary series with their final value.
        up_q_m, low_z_m = {}, {}
        for sec in self.upper_boundary_sections_sorted:
            up_q_m[sec] = self._pad_1d_last(evt.upper_q[sec], Tm)
        for sec in self.lower_boundary_sections_sorted:
            low_z_m[sec] = self._pad_1d_last(evt.lower_z[sec], Tm)

        first_up = self.upper_boundary_sections_sorted[0]
        first_low = self.lower_boundary_sections_sorted[0]
        q_up = up_q_m[first_up]
        z_low = low_z_m[first_low]

        lateral_q = self._pad_2d_last(evt.lateral, Tm)
        lat_density, lat_cum = self._build_lateral_features(lateral_q)
        node_source = self._pad_2d_last(evt.node_source, Tm)

        h_low = z_low - self.section_z_bed[first_low]
        h0 = evt.z0 - self.z_bed_array

        if i_evt_for_conditioning is not None and i_evt_for_conditioning < len(self._evt_log_qmax):
            log_qmax = self._evt_log_qmax[i_evt_for_conditioning]
            log_hrange = self._evt_log_hrange[i_evt_for_conditioning]
            q_scale_evt = self._evt_q_scale[i_evt_for_conditioning]
            h_range_evt = self._evt_h_range[i_evt_for_conditioning]
        else:
            feat = self._event_scale_features(evt)
            log_qmax = feat['log_qmax']
            log_hrange = feat['log_hrange']
            q_scale_evt = feat['q_scale_evt']
            h_range_evt = feat['h_range_evt']

        grid = np.zeros((self.in_channels, N, Tm), dtype=np.float32)
        grid[self.CHANNELS['x_norm'], :, :] = self.x_norm[:, None]
        grid[self.CHANNELS['x_local_norm'], :, :] = self.x_local_norm[:, None]
        grid[self.CHANNELS['x_global_norm'], :, :] = self.x_global_norm[:, None]
        grid[self.CHANNELS['reach_norm'], :, :] = self.reach_norm[:, None]
        grid[self.CHANNELS['z_bed_norm'], :, :] = self.z_bed_norm[:, None]
        grid[self.CHANNELS['S0_norm'], :, :] = self.S0_norm[:, None]
        grid[self.CHANNELS['n_norm'], :, :] = self.n_norm[:, None]
        grid[self.CHANNELS['tau'], :, :] = tau_model[None, :]
        grid[self.CHANNELS['time_mask'], :, :] = time_mask_1d[None, :]
        grid[self.CHANNELS['dt_hours'], :, :] = float(dt_h)

        # Aggregate available upstream-Q and downstream-stage channels.
        grid[self.CHANNELS['q_up'], :, :] = self._norm_q(q_up)[None, :]
        grid[self.CHANNELS['h_low'], :, :] = self._norm_h(h_low)[None, :]
        grid[self.CHANNELS['q_up_evt'], :, :] = self._norm_q_evt(q_up, q_scale_evt)[None, :]
        grid[self.CHANNELS['h_low_rel'], :, :] = self._norm_h_rel_evt(h_low, h_range_evt)[None, :]

        # Use global and event-relative scales for sources and initial states.
        grid[self.CHANNELS['lateral'], :, :] = self._norm_source(lat_density)
        grid[self.CHANNELS['lateral_cum'], :, :] = self._norm_source(lat_cum)
        grid[self.CHANNELS['node_source'], :, :] = self._norm_source(node_source)
        grid[self.CHANNELS['lateral_evt'], :, :] = self._norm_source_evt(lat_density, q_scale_evt)
        grid[self.CHANNELS['lateral_cum_evt'], :, :] = self._norm_source_evt(lat_cum, q_scale_evt)
        grid[self.CHANNELS['node_source_evt'], :, :] = self._norm_source_evt(node_source, q_scale_evt)
        grid[self.CHANNELS['h0'], :, :] = self._norm_h(h0)[:, None]
        grid[self.CHANNELS['q0'], :, :] = self._norm_q(evt.q0)[:, None]
        grid[self.CHANNELS['q0_evt'], :, :] = self._norm_q_evt(evt.q0, q_scale_evt)[:, None]
        grid[self.CHANNELS['duration'], :, :] = (duration_h / self.duration_scale)

        grid[self.CHANNELS['log_qmax_evt'], :, :] = (log_qmax - self.log_qmax_center) / self.log_qmax_scale
        grid[self.CHANNELS['log_hrange_evt'], :, :] = (log_hrange - self.log_hrange_center) / self.log_hrange_scale

        # Populate one pair of channels per available external boundary.
        for sec in self.upper_boundary_sections_sorted:
            grid[self.CHANNELS[f'up_{sec}_q'], :, :] = self._norm_q(up_q_m[sec])[None, :]
            grid[self.CHANNELS[f'up_{sec}_q_evt'], :, :] = self._norm_q_evt(up_q_m[sec], q_scale_evt)[None, :]
            rows = self.row_indices_by_section.get(sec, [])
            if rows:
                grid[self.CHANNELS[f'mask_up_{sec}'], rows, :] = 1.0
        for sec in self.lower_boundary_sections_sorted:
            bed = self.section_z_bed[sec]
            h_sec = low_z_m[sec] - bed
            grid[self.CHANNELS[f'low_{sec}_h'], :, :] = self._norm_h(h_sec)[None, :]
            grid[self.CHANNELS[f'low_{sec}_h_rel'], :, :] = self._norm_h_rel_evt(h_sec, h_range_evt)[None, :]
            rows = self.row_indices_by_section.get(sec, [])
            if rows:
                grid[self.CHANNELS[f'mask_low_{sec}'], rows, :] = 1.0

        return torch.tensor(grid, dtype=torch.float32, device=self.device), duration_h, time_mask_1d, valid_T, dt_h

    def _build_prepared_events(self):
        self.events.clear()
        for i_evt, evt in enumerate(self.raw_events):
            input_grid, duration_h, time_mask_1d, valid_T, dt_h = self._build_input_grid_for_event(
                evt, i_evt_for_conditioning=i_evt)

            target_grid = None
            if evt.target_z is not None and evt.target_q is not None:
                z_tar = self._pad_2d_last(evt.target_z, self.n_time_model)
                q_tar = self._pad_2d_last(evt.target_q, self.n_time_model)
                h_tar = z_tar - self.z_bed_array[:, None]
                target = np.stack([self._norm_h(h_tar), self._norm_q(q_tar)], axis=0).astype(np.float32)
                target_grid = torch.tensor(target, dtype=torch.float32, device=self.device)

            reference_grid = None
            if evt.reference_z is not None and evt.reference_q is not None:
                z_ref = self._pad_2d_last(evt.reference_z, self.n_time_model)
                q_ref = self._pad_2d_last(evt.reference_q, self.n_time_model)
                h_ref = z_ref - self.z_bed_array[:, None]
                ref = np.stack([self._norm_h(h_ref), self._norm_q(q_ref)], axis=0).astype(np.float32)
                reference_grid = torch.tensor(ref, dtype=torch.float32, device=self.device)

            obs_mask = np.zeros((len(self.row_sec_ids), self.n_time_model), dtype=np.float32)
            if evt.observed_section_indices is None:
                # None selects all available target rows.
                if target_grid is not None:
                    obs_mask[:, :valid_T] = 1.0
            elif len(evt.observed_section_indices) == 0:
                # [] keeps targets for evaluation but excludes them from training.
                pass
            else:
                # A list selects sparse supervised sections.
                for idx in evt.observed_section_indices:
                    obs_mask[idx, :valid_T] = 1.0
            reference_mask = np.zeros((len(self.row_sec_ids), self.n_time_model), dtype=np.float32)
            if evt.reference_section_indices is not None and reference_grid is not None:
                for idx in evt.reference_section_indices:
                    reference_mask[idx, :valid_T] = 1.0

            obs_mask_tensor = torch.tensor(obs_mask, dtype=torch.float32, device=self.device)
            reference_mask_tensor = torch.tensor(reference_mask, dtype=torch.float32, device=self.device)
            time_mask_tensor = torch.tensor(time_mask_1d, dtype=torch.float32, device=self.device)

            self.events.append(PreparedEvent(
                name=evt.name, t_hours_raw=evt.t_hours, duration_hours=duration_h,
                input_grid=input_grid,
                target_grid=target_grid, obs_mask=obs_mask_tensor,
                reference_grid=reference_grid, reference_mask=reference_mask_tensor,
                time_mask=time_mask_tensor, valid_T=valid_T, dt_hours=dt_h))

    def _build_query_coords(self):
        """Build DeepONet query coordinates from physical topology metrics."""
        tau = np.linspace(0.0, 1.0, self.n_time_model, dtype=np.float32)
        n_nodes = len(self.row_sec_ids)
        n_reaches = len(self.reach_names)

        def log_distance_norm(values):
            values = np.asarray(values, dtype=float)
            finite = np.isfinite(values)
            fallback = float(np.max(values[finite])) if np.any(finite) else 0.0
            values = np.where(finite, values, fallback)
            transformed = np.log1p(np.maximum(values, 0.0))
            scale = max(float(np.max(transformed)), 1.0)
            return transformed / scale

        local_fraction = np.zeros(n_nodes, dtype=float)
        reach_length_norm = np.zeros(n_nodes, dtype=float)
        reach_one_hot = np.zeros((n_nodes, n_reaches), dtype=float)
        is_head = np.zeros(n_nodes, dtype=float)
        is_tail = np.zeros(n_nodes, dtype=float)
        for reach_idx, reach_name in enumerate(self.reach_names):
            rows = self.reach_row_indices[reach_name]
            length = max(float(self.reach_lengths[reach_name]), 1.0)
            local_fraction[rows] = self.x_local_array[rows] / length
            reach_length_norm[rows] = np.log1p(length)
            reach_one_hot[rows, reach_idx] = 1.0
        reach_length_norm /= max(float(reach_length_norm.max()), 1.0)
        for meta in self.row_to_meta:
            is_head[meta['row']] = float(meta['is_head'])
            is_tail[meta['row']] = float(meta['is_tail'])

        junction_degree = self.junction_degree_array / max(
            float(np.max(self.junction_degree_array)), 1.0)
        static_features = np.column_stack([
            local_fraction,
            log_distance_norm(self.distance_from_upstream_m),
            log_distance_norm(self.distance_to_downstream_m),
            log_distance_norm(self.distance_to_junction_m),
            reach_length_norm,
            self.z_bed_norm,
            self.S0_norm,
            self.n_norm,
            is_head,
            is_tail,
            junction_degree,
            reach_one_hot,
        ]).astype(np.float32)

        static_grid = np.repeat(
            static_features[:, None, :], self.n_time_model, axis=1)
        time_grid = np.repeat(
            tau[None, :, None], n_nodes, axis=0)
        coords = np.concatenate([static_grid, time_grid], axis=2)
        self.query_dim = int(coords.shape[2])
        self.query_coords = torch.tensor(
            coords.reshape(-1, self.query_dim),
            dtype=torch.float32,
            device=self.device,
        )

    def _build_deeponet_model(self):
        """Create the DeepONet after constructing the river topology."""
        n_nodes = len(self.row_sec_ids)
        n_reaches = len(self.reach_names)
        reach_membership = torch.zeros(
            n_reaches, n_nodes, dtype=torch.float32, device=self.device)
        for reach_idx, reach_name in enumerate(self.reach_names):
            reach_membership[reach_idx, self.reach_row_indices[reach_name]] = 1.0

        junction_mask = torch.zeros(
            n_nodes, dtype=torch.float32, device=self.device)
        for node in self.junctions.values():
            junction_mask[node['rows']] = 1.0

        # Display-only global coordinates are excluded from the branch encoder.
        excluded = {'x_norm', 'x_global_norm', 'reach_norm'}
        branch_channels = [
            index for name, index in self.CHANNELS.items()
            if name not in excluded
        ]
        self.branch_channel_indices_tensor = torch.tensor(
            branch_channels, dtype=torch.long, device=self.device)
        selected_position = {
            original_index: position
            for position, original_index in enumerate(branch_channels)
        }
        boundary_channel_positions = torch.tensor([
            selected_position[self.CHANNELS[name]]
            for name in self.boundary_route_names
        ], dtype=torch.long, device=self.device)
        boundary_route_weights = torch.tensor(
            self.boundary_route_weights_array,
            dtype=torch.float32,
            device=self.device,
        )
        boundary_lag_fraction = torch.tensor(
            self.boundary_lag_fraction_array,
            dtype=torch.float32,
            device=self.device,
        )

        # Build physical neighboring-section and junction edges.
        edge_src, edge_dst, edge_dx, raw_attrs = [], [], [], []

        def append_edge(src, dst, dx_m, slope, roughness, direction, junction):
            edge_src.append(int(src))
            edge_dst.append(int(dst))
            edge_dx.append(float(dx_m))
            raw_attrs.append([
                float(dx_m),
                float(slope),
                float(roughness),
                float(direction),
                float(junction),
            ])

        for reach_name in self.reach_names:
            rows = self.reach_row_indices[reach_name]
            local_x = self.reach_local_x[reach_name]
            for i in range(len(rows) - 1):
                u, v = int(rows[i]), int(rows[i + 1])
                dx_m = float(local_x[i + 1] - local_x[i])
                slope = 0.5 * (
                    float(self.S0_array[u]) + float(self.S0_array[v]))
                roughness = 0.5 * (
                    float(self.manning_n_array[u])
                    + float(self.manning_n_array[v]))
                append_edge(
                    u, v, dx_m, slope, roughness,
                    direction=1.0, junction=0.0)
                append_edge(
                    v, u, dx_m, slope, roughness,
                    direction=-1.0, junction=0.0)

        for node in self.junctions.values():
            rows = [int(r) for r in node['rows']]
            upstream_rows = set(int(r) for r in node['upstream_rows'])
            downstream_rows = set(int(r) for r in node['downstream_rows'])
            for src in rows:
                for dst in rows:
                    if src == dst:
                        continue
                    if src in upstream_rows and dst in downstream_rows:
                        direction = 1.0
                    elif src in downstream_rows and dst in upstream_rows:
                        direction = -1.0
                    else:
                        direction = 0.0
                    roughness = 0.5 * (
                        float(self.manning_n_array[src])
                        + float(self.manning_n_array[dst]))
                    append_edge(
                        src, dst, 0.0, 0.0, roughness,
                        direction=direction, junction=1.0)

        edge_index = torch.tensor(
            [edge_src, edge_dst], dtype=torch.long, device=self.device)
        edge_dx_m = torch.tensor(
            edge_dx, dtype=torch.float32, device=self.device)
        raw_attrs = np.asarray(raw_attrs, dtype=np.float32)
        if raw_attrs.size == 0:
            edge_attr_array = np.zeros((0, 5), dtype=np.float32)
        else:
            dx_feature = np.log1p(np.maximum(raw_attrs[:, 0], 0.0))
            dx_feature /= max(float(dx_feature.max()), 1.0)
            slope_scale = max(
                float(np.quantile(np.abs(raw_attrs[:, 1]), 0.95)), 1e-6)
            slope_feature = np.tanh(raw_attrs[:, 1] / slope_scale)
            roughness_scale = max(float(raw_attrs[:, 2].max()), 0.01)
            roughness_feature = raw_attrs[:, 2] / roughness_scale
            edge_attr_array = np.column_stack([
                dx_feature,
                slope_feature,
                roughness_feature,
                raw_attrs[:, 3],
                raw_attrs[:, 4],
            ]).astype(np.float32)
        edge_attr = torch.tensor(
            edge_attr_array, dtype=torch.float32, device=self.device)
        dt_channel_position = selected_position[
            self.CHANNELS['dt_hours']]

        self.model = TopologyDeepONet(
            in_channels=len(branch_channels),
            query_dim=self.query_dim,
            hidden=self.deeponet_hidden,
            latent=self.deeponet_latent,
            reach_membership=reach_membership,
            junction_mask=junction_mask,
            boundary_channel_positions=boundary_channel_positions,
            boundary_route_weights=boundary_route_weights,
            boundary_lag_fraction=boundary_lag_fraction,
            dt_channel_position=dt_channel_position,
            edge_index=edge_index,
            edge_attr=edge_attr,
            edge_dx_m=edge_dx_m,
        ).to(self.device)
        print(
            f"DeepONet built: branch channels={len(branch_channels)}, "
            f"trunk dimension={self.query_dim}, reaches={n_reaches}, "
            f"boundary series={len(self.boundary_route_names)}, "
            f"graph edges={edge_index.shape[1]}",
            flush=True,
        )

    # ----------------------------------------------------------
    # Forward pass, hard constraints, and losses
    # ----------------------------------------------------------

    def _split_indices(self):
        name_to_idx = {e.name: i for i, e in enumerate(self.events)}
        train_idx = [name_to_idx[n] for n in self.train_event_names if n in name_to_idx]
        val_idx = [name_to_idx[n] for n in self.val_event_names if n in name_to_idx]
        test_idx = [name_to_idx[n] for n in self.test_event_names if n in name_to_idx]
        return train_idx, val_idx, test_idx

    def _forward_batch(self, x):
        branch_input = x.index_select(
            1, self.branch_channel_indices_tensor)
        return self.model(
            branch_input,
            self.query_coords,
            n_nodes=len(self.row_sec_ids),
            n_times=self.n_time_model,
        )

    def _fixed_z_rows(self):
        rows = []
        for sec in self.lower_boundary_sections_sorted:
            rows.extend(self.row_indices_by_section.get(sec, []))
        return set(rows)

    def _fixed_q_rows(self):
        rows = []
        for sec in self.upper_boundary_sections_sorted:
            rows.extend(self.row_indices_by_section.get(sec, []))
        return set(rows)

    def _apply_boundary_hard_injection(self, pred, x):
        """Apply initial, external-boundary, and junction hard constraints."""
        pred = pred.clone()

        # Initial conditions.
        pred[:, 0, :, 0] = x[:, self.CHANNELS['h0'], :, 0]
        pred[:, 1, :, 0] = x[:, self.CHANNELS['q0'], :, 0]

        # Prescribed discharge at upstream boundaries.
        for sec in self.upper_boundary_sections_sorted:
            rows = self.row_indices_by_section.get(sec, [])
            if not rows:
                continue
            ch = self.CHANNELS[f'up_{sec}_q']
            for r in rows:
                pred[:, 1, r, :] = x[:, ch, r, :]

        # Prescribed stage/depth at downstream boundaries.
        for sec in self.lower_boundary_sections_sorted:
            rows = self.row_indices_by_section.get(sec, [])
            if not rows:
                continue
            ch = self.CHANNELS[f'low_{sec}_h']
            for r in rows:
                pred[:, 0, r, :] = x[:, ch, r, :]

        if len(self.junctions) == 0:
            return pred

        # Project junction states in physical units, then normalize again.
        B_, _, N, T = pred.shape
        h = torch.clamp(self._denorm_h(pred[:, 0]), min=0.01)
        Z = h + self.z_bed_tensor.view(1, N, 1)
        Q = self._denorm_q(pred[:, 1])

        Z, Q = self._apply_junction_projection_physical(Z, Q, x)

        pred[:, 0] = self._norm_h(torch.clamp(Z - self.z_bed_tensor.view(1, N, 1), min=0.01))
        pred[:, 1] = self._norm_q(Q)
        return pred

    def _apply_junction_projection_physical(self, Z, Q, x):
        fixed_z_rows_global = self._fixed_z_rows()
        fixed_q_rows_global = self._fixed_q_rows()
        eps = 1e-12

        for node_sec_id, node in self.junctions.items():
            endpoint_meta = node['endpoint_meta']
            rows = [m['row'] for m in endpoint_meta]

            # Enforce stage continuity at each junction.
            z_fixed_rows = [r for r in rows if r in fixed_z_rows_global]
            if z_fixed_rows:
                z_target = Z[:, z_fixed_rows[0]:z_fixed_rows[0] + 1, :]
            else:
                z_target = Z[:, rows, :].mean(dim=1, keepdim=True)
            for r in rows:
                if (r not in z_fixed_rows) or (not z_fixed_rows):
                    Z[:, r:r + 1, :] = z_target

            # Preserve predicted branch ratios while enforcing total discharge.
            ordered_rows, signs, q_fixed_rows = [], [], []
            for m in endpoint_meta:
                row = m['row']
                ordered_rows.append(row)
                signs.append(1.0 if m['is_head'] else -1.0)
                if row in fixed_q_rows_global:
                    q_fixed_rows.append(row)

            signs_t = torch.tensor(signs, dtype=torch.float32, device=self.device).view(1, -1, 1)
            q_vals = Q[:, ordered_rows, :]
            fixed_mask = torch.tensor([1.0 if r in q_fixed_rows else 0.0 for r in ordered_rows],
                                      dtype=torch.float32, device=self.device).view(1, -1, 1)
            free_mask = 1.0 - fixed_mask
            denom = torch.sum((signs_t ** 2) * free_mask, dim=1, keepdim=True)

            # Every occurrence carries the same node-source value.
            node_src_norm = x[:, self.CHANNELS['node_source'], rows[0], :]
            node_src = self._denorm_source(node_src_norm).view(x.shape[0], 1, x.shape[-1])
            upstream_pos = [
                j for j, m in enumerate(endpoint_meta) if m['is_tail']]
            downstream_pos = [
                j for j, m in enumerate(endpoint_meta) if m['is_head']]

            if downstream_pos and not any(
                    ordered_rows[j] in q_fixed_rows
                    for j in downstream_pos):
                q_up_total = q_vals[:, upstream_pos, :].sum(
                    dim=1, keepdim=True)
                q_down_target = q_up_total + node_src
                q_new = q_vals.clone()
                if len(downstream_pos) == 1:
                    q_new[
                        :, downstream_pos[0]:downstream_pos[0] + 1, :
                    ] = q_down_target
                else:
                    q_down_raw = q_vals[:, downstream_pos, :]
                    weights = F.softplus(
                        q_down_raw / max(float(self.q_scale), 1e-6))
                    weights = weights / weights.sum(
                        dim=1, keepdim=True).clamp(min=eps)
                    q_new[:, downstream_pos, :] = (
                        weights * q_down_target)
            else:
                residual = (
                    torch.sum(signs_t * q_vals, dim=1, keepdim=True)
                    - node_src)
                if torch.any(denom > eps):
                    correction = residual / torch.clamp(denom, min=eps)
                    q_new = (
                        q_vals - signs_t * correction * free_mask)
                else:
                    q_new = q_vals

            for j, r in enumerate(ordered_rows):
                Q[:, r, :] = q_new[:, j, :]

        return Z, Q

    def _masked_data_loss_h(self, pred_h_norm, target_h_norm, mask):
        """Public water-level loss: standard masked MSE."""
        return self._masked_mse(pred_h_norm, target_h_norm, mask)

    def _masked_data_loss_q(self, pred_q_norm, target_q_norm, mask):
        """Public discharge loss: standard masked MSE."""
        return self._masked_mse(pred_q_norm, target_q_norm, mask)

    @staticmethod
    def _masked_mse(pred, target, mask):
        """Compute MSE only at valid observation times and locations."""
        weight = mask.to(dtype=pred.dtype)
        denom = weight.sum().clamp(min=1.0)
        return (((pred - target) ** 2) * weight).sum() / denom

    def _masked_physical_rmse(self, pred, target, mask):
        """Compute validation RMSE in physical units."""
        n_obs = mask.sum().clamp(min=1.0)
        h_error = (
            self._denorm_h(pred[:, 0])
            - self._denorm_h(target[:, 0])
        ) * mask
        q_error = (
            self._denorm_q(pred[:, 1])
            - self._denorm_q(target[:, 1])
        ) * mask
        z_rmse = torch.sqrt((h_error.square() * mask).sum() / n_obs)
        q_rmse = torch.sqrt((q_error.square() * mask).sum() / n_obs)
        # Equal-weight dimensionless score used only for model selection.
        score = 0.5 * (
            z_rmse / max(float(self.h_scale), 1e-6)
            + q_rmse / max(float(self.q_scale), 1e-6)
        )
        return z_rmse, q_rmse, score

    def _boundary_loss(self, pred, x):
        """Masked MSE between predictions and prescribed boundaries."""
        loss = torch.tensor(0.0, device=self.device)
        n_terms = 0
        tmask = x[:, self.CHANNELS['time_mask'], 0, :]
        for sec in self.upper_boundary_sections_sorted:
            ch = self.CHANNELS[f'up_{sec}_q']
            for r in self.row_indices_by_section.get(sec, []):
                loss = loss + self._masked_mse(
                    pred[:, 1, r, :], x[:, ch, r, :], tmask)
                n_terms += 1
        for sec in self.lower_boundary_sections_sorted:
            ch = self.CHANNELS[f'low_{sec}_h']
            for r in self.row_indices_by_section.get(sec, []):
                loss = loss + self._masked_mse(
                    pred[:, 0, r, :], x[:, ch, r, :], tmask)
                n_terms += 1
        return loss / max(n_terms, 1)

    def _junction_loss(self, pred, x):
        """Disabled extension hook for a junction soft constraint."""
        return pred.sum() * 0.0

    def _masked_peak_loss(self, pred, target, mask, peak_weight=5.0):
        """Disabled extension hook for peak-weighted supervision."""
        return pred.sum() * 0.0

    def _public_consistency_loss(self, pred, x):
        """
        Public first-order spatiotemporal consistency regularizer.

        This replaceable baseline uses centered differences of normalized state
        variables. It does not include the paper-specific equation residual,
        hydraulic/friction terms, or adaptive scaling strategy.
        """
        _, _, _, n_times = pred.shape
        zero = pred.sum() * 0.0
        if n_times < 2:
            return zero, zero

        eps = 1.0e-6
        valid_t = x[:, self.CHANNELS['time_mask'], 0, :]
        valid_pair = valid_t[:, :-1] * valid_t[:, 1:]
        tau = x[:, self.CHANNELS['tau'], 0, :]
        dtau = (tau[:, 1:] - tau[:, :-1]).abs().clamp(min=eps)

        total_cont = zero
        total_mom = zero
        total_weight = zero
        for reach_name in self.reach_names:
            rows = self.reach_row_indices[reach_name]
            if len(rows) < 2:
                continue

            h = pred[:, 0, rows, :]
            q = pred[:, 1, rows, :]
            dh_dt = (h[..., 1:] - h[..., :-1]) / dtau[:, None, :]
            dq_dt = (q[..., 1:] - q[..., :-1]) / dtau[:, None, :]

            x_local = x[:, self.CHANNELS['x_local_norm'], rows, 0]
            dx = (x_local[:, 1:] - x_local[:, :-1]).abs().clamp(min=eps)
            dh_dx = 0.5 * (
                h[:, 1:, 1:] - h[:, :-1, 1:]
                + h[:, 1:, :-1] - h[:, :-1, :-1]
            ) / dx[:, :, None]
            dq_dx = 0.5 * (
                q[:, 1:, 1:] - q[:, :-1, 1:]
                + q[:, 1:, :-1] - q[:, :-1, :-1]
            ) / dx[:, :, None]

            dh_dt_cell = 0.5 * (dh_dt[:, 1:, :] + dh_dt[:, :-1, :])
            dq_dt_cell = 0.5 * (dq_dt[:, 1:, :] + dq_dt[:, :-1, :])
            mask_cell = valid_pair[:, None, :].expand_as(dh_dx)

            # Public baseline: simple linear coupling terms.
            r_cont = dh_dt_cell + dq_dx
            r_mom = dq_dt_cell + dh_dx
            total_cont = total_cont + (r_cont.square() * mask_cell).sum()
            total_mom = total_mom + (r_mom.square() * mask_cell).sum()
            total_weight = total_weight + mask_cell.sum()

        denom = total_weight.clamp(min=1.0)
        return total_cont / denom, total_mom / denom

    # ----------------------------------------------------------
    # Training
    # ----------------------------------------------------------

    def train_model(self, epochs=3000, batch_size=2, lr=2e-3, lr_min=1e-5,
                    lambda_h=1.0, lambda_q=2.0,
                    lambda_pde_cont=5.0, lambda_pde_mom=5.0,
                    lambda_bc=0.5, lambda_peak=0.0, lambda_junction=0.0,
                    peak_weight=5.0, weight_decay=1e-5,
                    warmup_epochs=200, plateau_patience=200, plateau_factor=0.5,
                    pde_warmup_epochs=300,
                    early_stop_patience: Optional[int] = None,
                    verbose_every=50, diagnostics_file=None,
                    val_metric_mode='reference_data', val_include_pde=False):
        if not self.prepared:
            raise RuntimeError("Call prepare_data() before training")

        train_idx, val_idx, _ = self._split_indices()
        if not train_idx:
            raise RuntimeError("The training split is empty")

        train_loader = TorchDataLoader(
            EventTensorDataset(self.events, train_idx),
            batch_size=batch_size, shuffle=True)
        val_loader = (TorchDataLoader(EventTensorDataset(self.events, val_idx),
                                      batch_size=batch_size, shuffle=False)
                      if val_idx else None)

        optimizer = AdamW(self.model.parameters(), lr=lr, weight_decay=weight_decay)
        warmup_scheduler = torch.optim.lr_scheduler.LinearLR(
            optimizer, start_factor=0.01, end_factor=1.0, total_iters=max(warmup_epochs, 1))
        plateau_scheduler = torch.optim.lr_scheduler.ReduceLROnPlateau(
            optimizer, mode='min', factor=plateau_factor, patience=plateau_patience,
            min_lr=lr_min)

        best_state = None
        best_val = float('inf')
        bad_epochs = 0
        min_delta = 1e-8
        start_time = time.time()
        self.history = TrainingHistory()

        print("\n" + "=" * 72, flush=True)
        print(
            "DeepONet-PINO v8.1 public baseline: node data and first-order "
            "consistency regularization", flush=True)
        print(f"   λ_data: h={lambda_h}, q={lambda_q}, peak={lambda_peak}", flush=True)
        print(f"   λ_pde: cont={lambda_pde_cont}, mom={lambda_pde_mom}; λ_bc={lambda_bc}; λ_junction={lambda_junction}", flush=True)
        print(f"   warmup={warmup_epochs}ep", flush=True)
        print(f"   norm=affine(Q/h/source), data/bc=masked-MSE; overbank={self.overbank_mode}", flush=True)
        print(
            f"   Validation: prediction-only inputs; mode={val_metric_mode}; "
            f"include_consistency={val_include_pde}", flush=True)
        if early_stop_patience is not None:
            print(f"   early stopping patience={early_stop_patience} epochs", flush=True)
        print("=" * 72, flush=True)

        for ep in range(1, epochs + 1):
            self.model.train()
            s_loss = s_h = s_q = s_bc = s_pk = s_pc = s_pm = s_junc = 0.0
            n_batches = 0
            pde_factor = min(
                1.0,
                max(float(ep) / max(float(pde_warmup_epochs), 1.0), 0.05),
            )
            for batch in train_loader:
                x = batch['x'].to(self.device)
                y = batch['y'].to(self.device)
                mask = batch['mask'].to(self.device)
                optimizer.zero_grad()
                pred = self._forward_batch(x)
                if self.boundary_hard_injection:
                    pred = self._apply_boundary_hard_injection(pred, x)

                has_obs = mask.sum() > 0
                if has_obs:
                    loss_h = self._masked_data_loss_h(pred[:, 0], y[:, 0], mask)
                    loss_q = self._masked_data_loss_q(pred[:, 1], y[:, 1], mask)
                    loss_peak = self._masked_peak_loss(pred, y, mask, peak_weight)
                else:
                    loss_h = loss_q = loss_peak = torch.tensor(0.0, device=self.device)

                loss_pde_c, loss_pde_m = self._public_consistency_loss(pred, x)
                loss_bc = self._boundary_loss(pred, x)
                loss_junction = self._junction_loss(pred, x) if lambda_junction > 0 else torch.tensor(0.0, device=self.device)

                loss = (lambda_h * loss_h + lambda_q * loss_q
                        + pde_factor * (
                            lambda_pde_cont * loss_pde_c
                            + lambda_pde_mom * loss_pde_m)
                        + lambda_bc * loss_bc + lambda_peak * loss_peak
                        + lambda_junction * loss_junction)
                loss.backward()
                torch.nn.utils.clip_grad_norm_(self.model.parameters(), max_norm=1.0)
                optimizer.step()

                s_loss += loss.item(); s_h += loss_h.item(); s_q += loss_q.item()
                s_bc += loss_bc.item(); s_pk += loss_peak.item()
                s_pc += loss_pde_c.item(); s_pm += loss_pde_m.item(); s_junc += loss_junction.item()
                n_batches += 1

            nb = max(n_batches, 1)
            epoch_loss = s_loss / nb

            val_loss = float('nan')
            val_data_loss_print = float('nan')
            val_pde_loss_print = float('nan')
            val_z_rmse_print = float('nan')
            val_q_rmse_print = float('nan')
            if val_loader is not None:
                self.model.eval()
                total_v = 0.0
                total_v_data = 0.0
                total_v_pde = 0.0
                total_z_rmse = 0.0
                total_q_rmse = 0.0
                n_val = 0
                n_val_data = 0
                with torch.no_grad():
                    for batch in val_loader:
                        x = batch['x'].to(self.device)
                        pred = self._forward_batch(x)
                        if self.boundary_hard_injection:
                            pred = self._apply_boundary_hard_injection(pred, x)

                        # Validation references are never model inputs.
                        ref_y = batch['ref_y'].to(self.device)
                        ref_mask = batch['ref_mask'].to(self.device)
                        if ref_mask.sum() > 0:
                            v_h_ref = self._masked_data_loss_h(pred[:, 0], ref_y[:, 0], ref_mask)
                            v_q_ref = self._masked_data_loss_q(pred[:, 1], ref_y[:, 1], ref_mask)
                            v_data = lambda_h * v_h_ref + lambda_q * v_q_ref
                            v_z_rmse, v_q_rmse, v_physical_score = (
                                self._masked_physical_rmse(
                                    pred, ref_y, ref_mask)
                            )
                            total_v_data += v_data.item()
                            total_z_rmse += v_z_rmse.item()
                            total_q_rmse += v_q_rmse.item()
                            n_val_data += 1
                        else:
                            v_data = torch.tensor(0.0, device=self.device)
                            v_physical_score = torch.tensor(
                                0.0, device=self.device)

                        v_c, v_m = self._public_consistency_loss(pred, x)
                        v_pde = lambda_pde_cont * v_c + lambda_pde_mom * v_m
                        total_v_pde += v_pde.item()

                        if str(val_metric_mode).lower() in ['reference_data', 'data', 'ref'] and ref_mask.sum() > 0:
                            # Select the best model primarily by physical RMSE.
                            v_metric = (
                                v_physical_score
                                + (v_pde if val_include_pde else 0.0)
                            )
                        else:
                            # Fall back to consistency loss without references.
                            v_metric = v_pde

                        total_v += float(v_metric.item() if torch.is_tensor(v_metric) else v_metric)
                        n_val += 1

                val_loss = total_v / max(n_val, 1)
                val_data_loss_print = total_v_data / max(n_val_data, 1) if n_val_data > 0 else float('nan')
                val_pde_loss_print = total_v_pde / max(n_val, 1)
                val_z_rmse_print = (
                    total_z_rmse / max(n_val_data, 1)
                    if n_val_data > 0 else float('nan')
                )
                val_q_rmse_print = (
                    total_q_rmse / max(n_val_data, 1)
                    if n_val_data > 0 else float('nan')
                )
                metric = float(val_loss)
            else:
                metric = epoch_loss

            if ep <= warmup_epochs:
                warmup_scheduler.step()
            else:
                # Schedule by validation loss, or training loss without validation.
                plateau_scheduler.step(metric)

            if metric < best_val - min_delta:
                best_val = metric
                best_state = {k: v.detach().cpu().clone() for k, v in self.model.state_dict().items()}
                bad_epochs = 0
            else:
                bad_epochs += 1

            elapsed = time.time() - start_time
            self.history.epoch.append(ep); self.history.lr.append(optimizer.param_groups[0]['lr'])
            self.history.train_loss.append(epoch_loss); self.history.train_h.append(s_h/nb)
            self.history.train_q.append(s_q/nb); self.history.train_bc.append(s_bc/nb)
            self.history.train_peak.append(s_pk/nb)
            self.history.train_pde_cont.append(s_pc/nb); self.history.train_pde_mom.append(s_pm/nb)
            self.history.train_rel.append(s_junc/nb)
            self.history.val_loss.append(val_loss if isinstance(val_loss, float) else val_loss.item())
            self.history.elapsed_sec.append(elapsed)

            if ep % verbose_every == 0 or ep == 1:
                lr_c = optimizer.param_groups[0]['lr']
                msg = (f"Ep {ep:5d} | {elapsed:7.1f}s | LR={lr_c:.5g} | "
                       f"h={s_h/nb:.5f} q={s_q/nb:.5f} "
                       f"pde_c={s_pc/nb:.5f} pde_m={s_pm/nb:.5f} "
                       f"pde_w={pde_factor:.2f} "
                    #    f"bc={s_bc/nb:.5f} junc={s_junc/nb:.5f} pk={s_pk/nb:.5f} | "
                       f"pk={s_pk/nb:.5f} | "
                       f"total={epoch_loss:.5f}")
                if val_loader:
                    msg += f" val={float(val_loss):.5f}"
                    if np.isfinite(val_data_loss_print):
                        msg += f" val_data={val_data_loss_print:.5f}"
                    if np.isfinite(val_pde_loss_print):
                        msg += f" val_pde={val_pde_loss_print:.5f}"
                    if np.isfinite(val_z_rmse_print):
                        msg += f" Zrmse={val_z_rmse_print:.3f}m"
                    if np.isfinite(val_q_rmse_print):
                        msg += f" Qrmse={val_q_rmse_print:.2f}m3/s"
                    msg += f" best={best_val:.5f} bad={bad_epochs}"
                print(msg, flush=True)

            if early_stop_patience is not None and val_loader is not None and bad_epochs >= early_stop_patience:
                print(
                    f"Early stopping at epoch {ep} after {bad_epochs} "
                    "epochs without validation improvement", flush=True)
                break

        if best_state is not None:
            self.model.load_state_dict(best_state)
            print(
                f"Training complete; restored best weights, "
                f"best loss={best_val:.6f}", flush=True)

        if diagnostics_file:
            self._save_diagnostics(diagnostics_file)

    # ----------------------------------------------------------
    # Diagnostics and result export
    # ----------------------------------------------------------

    def _infer_event(self, evt):
        self.model.eval()
        with torch.no_grad():
            x = evt.input_grid.unsqueeze(0)
            pred = self._forward_batch(x)
            if self.boundary_hard_injection:
                pred = self._apply_boundary_hard_injection(pred, x)
            pred_np = pred[0].cpu().numpy()
        h_pred = self._denorm_h(pred_np[0])
        q_pred = self._denorm_q(pred_np[1])
        z_pred = h_pred + self.z_bed_array[:, None]
        return z_pred, q_pred

    def _load_full_result_for_event(self, spec):
        if spec.result_file is None or not os.path.exists(spec.result_file):
            return None, None, None
        return self._parse_result_file_sparse(spec.result_file, observed_sections_set=None)[:3]

    def _sheet_name_for_row(self, row):
        sec = self.row_sec_ids[row]
        meta = self.row_to_meta[row]
        if self.section_representative_row.get(sec, row) == row:
            name = str(sec)
        else:
            name = f"{sec}_{meta['reach_name']}_{meta['local_idx']}"
        invalid = ['\\', '/', '*', '?', ':', '[', ']']
        for ch in invalid:
            name = name.replace(ch, '_')
        return name[:31]

    def _wide_result_frame(self, t_values, values, value_prefix=''):
        """Convert a [node, time] array to a time-by-node table."""
        data = {'time_hours': np.round(np.asarray(t_values, dtype=float), 6)}
        used_names = set(data)
        for row in range(len(self.row_sec_ids)):
            base = self._sheet_name_for_row(row)
            name = f'{value_prefix}{base}'
            if name in used_names:
                suffix = 2
                while f'{name}_{suffix}' in used_names:
                    suffix += 1
                name = f'{name}_{suffix}'
            used_names.add(name)
            data[name] = np.round(np.asarray(values[row], dtype=float), 6)
        return pd.DataFrame(data)

    def save_training_results(
            self, output_dir='training_results', full_compare=True):
        """Export summary metrics and long-form predictions for each event."""
        if not self.prepared:
            raise RuntimeError("Prepare data and train the model first")
        os.makedirs(output_dir, exist_ok=True)
        self.model.eval()
        all_summaries = []
        for evt, spec in zip(self.events, self.event_specs):
            z_pred_model, q_pred_model = self._infer_event(evt)
            t_raw = evt.t_hours_raw
            valid_T = min(evt.valid_T, len(t_raw), z_pred_model.shape[1])
            N = len(self.row_sec_ids)
            # Padding beyond the original event duration is not exported.
            z_pred = z_pred_model[:, :valid_T]
            q_pred = q_pred_model[:, :valid_T]
            t_raw = t_raw[:valid_T]

            z_true_full, q_true_full = None, None
            has_full = np.zeros(N, dtype=bool)
            if full_compare:
                t_full, zf, qf = self._load_full_result_for_event(spec)
                if t_full is not None and zf is not None:
                    z_true_full = np.zeros((N, len(t_raw)))
                    q_true_full = np.zeros((N, len(t_raw)))
                    for i in range(N):
                        if not np.all(np.isnan(zf[i])):
                            z_true_full[i] = np.interp(t_raw, t_full, zf[i])
                            q_true_full[i] = np.interp(t_raw, t_full, qf[i])
                            has_full[i] = True

            mask_np = evt.obs_mask.cpu().numpy()
            out_file = os.path.join(output_dir, f'{evt.name}.xlsx')
            summary_rows = []
            result_rows = []

            for i, sec_id in enumerate(self.row_sec_ids):
                meta = self.row_to_meta[i]
                node_name = self._sheet_name_for_row(i)
                is_obs = mask_np[i, :].sum() > 0
                has_true = has_full[i] if z_true_full is not None else False

                row = {
                    'row': i, 'section_id': sec_id,
                    'reach': meta['reach_name'],
                    'node_name': node_name,
                    'local_index': meta['local_idx'],
                    'source': 'observed' if is_obs else 'model',
                    'bed_elevation_m': round(self.z_bed_array[i], 3),
                }
                if has_true:
                    z_error = z_pred[i] - z_true_full[i]
                    q_error = q_pred[i] - q_true_full[i]
                    z_rmse = np.sqrt(np.mean(z_error ** 2))
                    q_rmse = np.sqrt(np.mean(q_error ** 2))
                    z_var = np.sum((z_true_full[i] - z_true_full[i].mean()) ** 2)
                    q_var = np.sum((q_true_full[i] - q_true_full[i].mean()) ** 2)
                    row.update({
                        'Z_RMSE_m': round(z_rmse, 4),
                        'Q_RMSE_m3s': round(q_rmse, 4),
                        'Z_NSE': round(1 - np.sum(z_error ** 2) / max(z_var, 1e-10), 4),
                        'Q_NSE': round(1 - np.sum(q_error ** 2) / max(q_var, 1e-10), 4),
                        'comparison_points': int(len(t_raw)),
                    })
                else:
                    row.update({
                        'Z_RMSE_m': np.nan,
                        'Q_RMSE_m3s': np.nan,
                        'Z_NSE': np.nan,
                        'Q_NSE': np.nan,
                        'comparison_points': 0,
                    })
                summary_rows.append(row)

                for j, time_h in enumerate(t_raw):
                    detail = {
                        'node_name': node_name,
                        'section_id': int(sec_id),
                        'reach': meta['reach_name'],
                        'local_index': int(meta['local_idx']),
                        'time_hours': round(float(time_h), 6),
                        'predicted_water_level_m': round(
                            float(z_pred[i, j]), 6),
                        'predicted_water_depth_m': round(
                            float(z_pred[i, j] - self.z_bed_array[i]), 6),
                        'predicted_discharge_m3s': round(
                            float(q_pred[i, j]), 6),
                        'source': 'observed' if is_obs else 'model',
                    }
                    if has_true:
                        detail.update({
                            'reference_water_level_m': round(
                                float(z_true_full[i, j]), 6),
                            'water_level_error_m': round(
                                float(z_pred[i, j] - z_true_full[i, j]), 6),
                            'reference_discharge_m3s': round(
                                float(q_true_full[i, j]), 6),
                            'discharge_error_m3s': round(
                                float(q_pred[i, j] - q_true_full[i, j]), 6),
                        })
                    else:
                        detail.update({
                            'reference_water_level_m': np.nan,
                            'water_level_error_m': np.nan,
                            'reference_discharge_m3s': np.nan,
                            'discharge_error_m3s': np.nan,
                        })
                    result_rows.append(detail)

            df_summary = pd.DataFrame(summary_rows)
            summary_columns = [
                'row', 'node_name', 'section_id', 'reach', 'local_index',
                'source', 'bed_elevation_m',
                'Z_RMSE_m', 'Z_NSE', 'Q_RMSE_m3s', 'Q_NSE',
                'comparison_points',
            ]
            df_summary = df_summary.reindex(columns=summary_columns)
            df_results = pd.DataFrame(result_rows)
            result_columns = [
                'node_name', 'section_id', 'reach', 'local_index', 'time_hours',
                'predicted_water_level_m', 'reference_water_level_m',
                'water_level_error_m', 'predicted_water_depth_m',
                'predicted_discharge_m3s', 'reference_discharge_m3s',
                'discharge_error_m3s', 'source',
            ]
            df_results = df_results.reindex(columns=result_columns)

            with pd.ExcelWriter(out_file, engine='openpyxl') as writer:
                df_summary.to_excel(writer, sheet_name='summary', index=False)
                df_results.to_excel(
                    writer, sheet_name='predictions', index=False)

            # Apply basic workbook formatting.
            try:
                import openpyxl
                from openpyxl.styles import Alignment, Font, PatternFill
                from openpyxl.utils import get_column_letter

                wb = openpyxl.load_workbook(out_file)
                header_fill = PatternFill(
                    fill_type='solid', fgColor='1F4E78')
                header_font = Font(color='FFFFFF', bold=True)
                for ws in (wb['summary'], wb['predictions']):
                    ws.freeze_panes = 'A2'
                    ws.auto_filter.ref = ws.dimensions
                    ws.sheet_view.showGridLines = False
                    for cell in ws[1]:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(
                            horizontal='center', vertical='center')
                    ws.row_dimensions[1].height = 24
                    for column_cells in ws.iter_cols(
                            min_row=1, max_row=min(ws.max_row, 300)):
                        max_len = max(
                            len(str(cell.value)) if cell.value is not None else 0
                            for cell in column_cells)
                        width = min(max(max_len + 2, 10), 28)
                        ws.column_dimensions[
                            get_column_letter(column_cells[0].column)
                        ].width = width
                wb.save(out_file)
                wb.close()
            except Exception as exc:
                warnings.warn(f"Could not format the result workbook: {exc}")

            print(f"Exported {evt.name} to {out_file}", flush=True)
            all_summaries.append({'event': evt.name, 'summary': df_summary})

        print(f"Training-domain results exported to: {output_dir}/", flush=True)
        return all_summaries

    def _save_diagnostics(self, file_path):
        writer = pd.ExcelWriter(file_path, engine='openpyxl')
        self.history.to_dataframe().to_excel(
            writer, sheet_name='training_history', index=False)
        topo_rows = []
        for m in self.row_to_meta:
            topo_rows.append({
                'row': m['row'], 'section_id': m['sec_id'],
                'reach': m['reach_name'],
                'local_idx': m['local_idx'], 'is_head': m['is_head'], 'is_tail': m['is_tail'],
                'x_local': m['x_local'], 'x_global': m['x_global'],
                'Z_bed': self.z_bed_array[m['row']]
            })
        pd.DataFrame(topo_rows).to_excel(
            writer, sheet_name='network_nodes', index=False)
        junc_rows = []
        for node_id, node in self.junctions.items():
            junc_rows.append({
                'junction_section': node_id,
                'endpoint_rows': ','.join(map(str, node['rows'])),
                'upstream_rows': ','.join(map(str, node['upstream_rows'])),
                'downstream_rows': ','.join(map(str, node['downstream_rows'])),
            })
        pd.DataFrame(junc_rows).to_excel(
            writer, sheet_name='junctions', index=False)
        writer.close()
        print(f"Diagnostics saved to: {file_path}", flush=True)

    # ----------------------------------------------------------
    # Prediction for new boundary conditions
    # ----------------------------------------------------------

    def predict_with_new_boundary(self, boundary_file, output_file,
                                  lateral_inflow_file=None, initial_condition_file=None):
        if not self.prepared:
            raise RuntimeError("Prepare data and train the model first")
        spec = EventSpec(name='predict', boundary_file=boundary_file, result_file=None,
                         lateral_inflow_file=lateral_inflow_file,
                         initial_condition_file=initial_condition_file)
        raw_evt = self._load_single_event(spec)
        old_T = self.n_time_model
        old_query_coords = self.query_coords
        use_extended_time = len(raw_evt.t_hours) > old_T
        if use_extended_time:
            print(
                f"Boundary length {len(raw_evt.t_hours)} exceeds training "
                f"max_T={old_T}; querying DeepONet at new time coordinates",
                flush=True,
            )
            self.n_time_model = len(raw_evt.t_hours)
            input_grid, duration_h, time_mask_1d, valid_T, dt_h = self._build_input_grid_for_event(raw_evt, i_evt_for_conditioning=None)
            self._build_query_coords()
        else:
            input_grid, duration_h, time_mask_1d, valid_T, dt_h = self._build_input_grid_for_event(raw_evt, i_evt_for_conditioning=None)

        try:
            self.model.eval()
            with torch.no_grad():
                x = input_grid.unsqueeze(0)
                pred = self._forward_batch(x)
                if self.boundary_hard_injection:
                    pred = self._apply_boundary_hard_injection(pred, x)
                pred_np = pred[0].cpu().numpy()
        finally:
            if use_extended_time:
                self.n_time_model = old_T
                self.query_coords = old_query_coords

        h_model = self._denorm_h(pred_np[0])
        z_model = h_model + self.z_bed_array[:, None]
        q_model = self._denorm_q(pred_np[1])

        N = len(self.row_sec_ids)
        out_T = min(valid_T, len(raw_evt.t_hours), z_model.shape[1])
        z_out = z_model[:, :out_T]
        q_out = q_model[:, :out_T]
        t_out = raw_evt.t_hours[:out_T]

        output_dir = os.path.dirname(output_file)
        if output_dir:
            os.makedirs(output_dir, exist_ok=True)
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            self._wide_result_frame(t_out, z_out).to_excel(
                writer, sheet_name='water_level', index=False)
            self._wide_result_frame(t_out, q_out).to_excel(
                writer, sheet_name='discharge', index=False)
        print(f"Prediction results saved to: {output_file}", flush=True)



def add_case_events(
    surrogate,
    root_dir='.',
    start_id=1,
    end_id=50,
    case_prefix='case_',
    event_prefix='event',
    observed_sections=None,
    boundary_filename='boundary_conditions.csv',
    lateral_filename='lateral_inflow.csv',
    initial_filename='initial_conditions.csv',
    result_filename='targets.csv',
    skip_missing=True,
):
    """
    Register a numbered collection of case directories.

    Existing manually registered events are preserved.

    Parameters
    ----------
    root_dir : str
        Parent directory containing numbered case folders.
    start_id, end_id : int
        Inclusive range of case identifiers.
    case_prefix : str
        Prefix before each zero-padded case identifier.
    event_prefix : str
        Prefix used for registered event names.
    observed_sections : None | [] | list[int]
        None selects full supervision, an empty list selects no observation
        supervision, and a list selects sparse sections.
    skip_missing : bool
        Skip incomplete cases when true; otherwise raise an error.

    Returns
    -------
    added_names : list[str]
        Names of successfully registered events.
    """
    root_dir = os.path.abspath(root_dir)
    added_names = []
    missing_cases = []

    print("\n" + "=" * 72, flush=True)
    print(f"Scanning case directory: {root_dir}", flush=True)
    print(f"   Requested IDs: {start_id} through {end_id}", flush=True)
    print("=" * 72, flush=True)

    for case_id in range(int(start_id), int(end_id) + 1):
        case_name = f"{case_prefix}{case_id:03d}"
        case_dir = os.path.join(root_dir, case_name)
        boundary_file = os.path.join(case_dir, boundary_filename)
        lateral_file = os.path.join(case_dir, lateral_filename)
        initial_file = os.path.join(case_dir, initial_filename)
        result_file = os.path.join(case_dir, result_filename)

        required = [boundary_file, lateral_file, initial_file, result_file]
        missing = [fp for fp in required if not os.path.exists(fp)]
        if missing:
            msg = f"{case_name} is missing: " + "; ".join(
                os.path.basename(x) for x in missing)
            if skip_missing:
                print(f"   Skipped: {msg}", flush=True)
                missing_cases.append((case_name, missing))
                continue
            raise FileNotFoundError(msg)

        event_name = f"{event_prefix}_{case_id:03d}"
        surrogate.add_event_sparse(
            name=event_name,
            boundary_file=boundary_file,
            lateral_inflow_file=lateral_file,
            initial_condition_file=initial_file,
            result_file=result_file,
            observed_sections=observed_sections,
        )
        added_names.append(event_name)

    print(
        f"Case registration complete: added={len(added_names)}, "
        f"skipped={len(missing_cases)}", flush=True)
    return added_names


def split_events(names, validation_case_ids):
    """Split registered event names into training and validation sets."""
    validation_case_ids = {int(x) for x in validation_case_ids}
    train_names, val_names, matched_ids = [], [], set()

    for name in names:
        try:
            case_id = int(str(name).rsplit('_', 1)[-1])
        except ValueError:
            case_id = -1
        if case_id in validation_case_ids:
            val_names.append(name)
            matched_ids.add(case_id)
        else:
            train_names.append(name)

    unmatched = sorted(validation_case_ids - matched_ids)
    if unmatched:
        print(f"Validation cases not found: {unmatched}", flush=True)
    return train_names, val_names


# ============================================================
# Command-line entry point
# ============================================================

if __name__ == '__main__':
    print("=" * 72, flush=True)
    print("Topology-aware river-network DeepONet baseline", flush=True)
    print("Node-level CSV input, physical reach distances, and DeepONet", flush=True)
    print("=" * 72, flush=True)

    repo_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    data_dir = os.path.join(repo_root, 'data')
    dl = DataLoader()
    dl.load_cross_sections(os.path.join(data_dir, 'cross_sections.csv'))
    dl.load_river_network(os.path.join(data_dir, 'river_network.csv'))
    dl.load_boundary_conditions(os.path.join(data_dir, 'boundary_sections.csv'))

    surrogate = RiverOperatorSurrogate(
        data_loader=dl, n_time_model=None,
        device=os.environ.get('DEVICE', 'cuda'),
        deeponet_hidden=128, deeponet_latent=128,
        boundary_hard_injection=True)

    # Optional manually registered training and validation events.
    manual_train_event_names = []
    manual_val_event_names = []

    # Bundled public-data cases.
    USE_CASES = True
    CASE_ROOT_DIR = os.environ.get('CASE_ROOT_DIR', data_dir)

    # All cases share one sparse-observation configuration.
    # Use None for full supervision, [] for no observations, or a list of nodes.
    OBSERVED_NODES = None

    CASE_START = 1
    CASE_END = 3

    VALIDATION_CASE_IDS = [3]

    case_event_names = []
    if USE_CASES:
        case_event_names = add_case_events(
            surrogate=surrogate,
            root_dir=CASE_ROOT_DIR,
            start_id=CASE_START,
            end_id=CASE_END,
            case_prefix=os.environ.get('PINN_CASE_PREFIX', 'case_'),
            event_prefix='event',
            observed_sections=OBSERVED_NODES,
            boundary_filename='boundary_conditions.csv',
            lateral_filename='lateral_inflow.csv',
            initial_filename='initial_conditions.csv',
            result_filename='targets.csv',
            skip_missing=True,
        )

    case_train_names, case_val_names = split_events(
        case_event_names, VALIDATION_CASE_IDS)

    surrogate.set_split(
        train_event_names=manual_train_event_names + case_train_names,
        val_event_names=manual_val_event_names + case_val_names,
        test_event_names=[]
    )

    # Validation is boundary-driven; targets are used only for evaluation.
    VAL_EVENTS_BOUNDARY_ONLY = True
    if VAL_EVENTS_BOUNDARY_ONLY:
        surrogate.set_events_observed_sections(
            manual_val_event_names + case_val_names, observed_sections=[])

    print("\nData split:", flush=True)
    print(f"   Observed nodes: {OBSERVED_NODES}", flush=True)
    print(f"   Cases: {len(case_event_names)}", flush=True)
    print(f"   Training cases: {len(case_train_names)}", flush=True)
    print(f"   Validation cases: {len(case_val_names)}", flush=True)
    print(f"   Validation IDs: {VALIDATION_CASE_IDS}", flush=True)

    surrogate.prepare_data()

    surrogate.train_model(
        epochs=int(os.environ.get('EPOCHS', '3000')),
        batch_size=int(os.environ.get('BATCH_SIZE', '1')),
        lr=1e-3, lr_min=1e-5,
        lambda_h=1.0, lambda_q=1.0,
        lambda_pde_cont=0.1, lambda_pde_mom=0.1,
        lambda_bc=0.0, lambda_peak=0.0, lambda_junction=0.0,
        warmup_epochs=200, plateau_patience=100, plateau_factor=0.5,
        early_stop_patience=500,
        diagnostics_file=os.path.join(repo_root, 'training_diagnostics.xlsx'),
        val_metric_mode='reference_data',
        val_include_pde=False)

    surrogate.save_training_results(
        output_dir=os.path.join(repo_root, 'training_results'),
        full_compare=True)

    prediction_dir = os.path.join(repo_root, 'prediction_case')
    prediction_boundary = os.path.join(
        prediction_dir, 'boundary_conditions.csv')
    if os.path.exists(prediction_boundary):
        surrogate.predict_with_new_boundary(
            boundary_file=prediction_boundary,
            lateral_inflow_file=os.path.join(
                prediction_dir, 'lateral_inflow.csv'),
            initial_condition_file=os.path.join(
                prediction_dir, 'initial_conditions.csv'),
            output_file=os.path.join(repo_root, 'prediction_results.xlsx'))
